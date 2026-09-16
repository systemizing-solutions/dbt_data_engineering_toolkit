"""Controlled workbook safety and registry-consistency checks."""

from __future__ import annotations

import yaml
from openpyxl import load_workbook

from .paths import ROOT
from dbt_data_engineering_toolkit_compiler.version import (
    COMPILER_VERSION,
    WORKBOOK_SCHEMA_VERSION,
)


def check_workbooks() -> list[str]:
    errors: list[str] = []
    compiler = ROOT / "python/src/dbt_data_engineering_toolkit_compiler"
    workbook = load_workbook(
        compiler / "resources/data_product.xlsx", read_only=False, data_only=False
    )
    required = {
        "Instructions",
        "Fundamentals",
        "Schema data_product",
        "DET Models",
        "DET Sources",
        "DET Source Schema",
        "DET Mapping",
        "DET Parameters",
        "DET Model Inputs",
        "DET Relationships",
        "Operational Validation",
        "Operational Parameters",
        "DET Lookups",
        "DET Build",
        "_DET Context",
        "_DET Lists",
        "_DET Metadata",
        "_DET Raw ODCS",
    }
    for missing in sorted(required - set(workbook.sheetnames)):
        errors.append(f"data_product.xlsx: missing required sheet {missing}")
    if not errors:
        quality_surfaces = {
            name
            for name in workbook.sheetnames
            if "quality" in name.casefold() or name.startswith("Operational ")
        }
        expected_quality_surfaces = {
            "Quality",
            "Operational Validation",
            "Operational Parameters",
        }
        if quality_surfaces != expected_quality_surfaces:
            errors.append(
                "data_product.xlsx: quality surfaces must be exactly Quality, "
                "Operational Validation, and Operational Parameters"
            )
        for name in ("_DET Context", "_DET Lists", "_DET Metadata", "_DET Raw ODCS"):
            if workbook[name].sheet_state != "veryHidden":
                errors.append(f"data_product.xlsx: {name} must be very hidden")
        if workbook["Quality"].sheet_state != "visible":
            errors.append(
                "data_product.xlsx: official Quality must be the visible authoring sheet"
            )
        if workbook["Quality"]["N4"].value != "Rule ID":
            errors.append("data_product.xlsx: official Quality is missing DET Rule ID")
        if workbook["Quality"]["P4"].value != "ODCS Passthrough (JSON)":
            errors.append(
                "data_product.xlsx: official Quality is missing ODCS passthrough"
            )
        for name in ("Fundamentals", "Schema data_product", "DET Mapping"):
            if not workbook[name].protection.sheet:
                errors.append(f"data_product.xlsx: {name} must be protected")
        if workbook["DET Mapping"].max_column != 6:
            errors.append("data_product.xlsx: DET Mapping must remain six columns")
        if workbook["DET Mapping"]["A4"].protection.locked:
            errors.append(
                "data_product.xlsx: business input cells must remain unlocked"
            )
        for name in (
            "DET Models",
            "DET Sources",
            "DET Source Schema",
            "DET Mapping",
            "DET Parameters",
            "Operational Validation",
            "Operational Parameters",
            "DET Lookups",
        ):
            if any(
                workbook[name].cell(4, column).value not in (None, "")
                for column in range(1, workbook[name].max_column + 1)
            ):
                errors.append(f"data_product.xlsx: {name} contains unsafe demo rows")
        if any(
            workbook["Quality"].cell(5, column).value not in (None, "")
            for column in range(1, 17)
        ):
            errors.append("data_product.xlsx: Quality contains unsafe demo rows")
        operators = yaml.safe_load(
            (compiler / "resources/operators.yml").read_text(encoding="utf-8")
        )["operators"]
        expected = {
            item["label"] for item in operators if item["kind"] == "transformation"
        }
        actual = {
            workbook["_DET Lists"].cell(row, 1).value
            for row in range(2, workbook["_DET Lists"].max_row + 1)
            if workbook["_DET Lists"].cell(row, 1).value
        }
        if actual != expected:
            errors.append("data_product.xlsx: dropdowns differ from operators.yml")
        if workbook["_DET Metadata"]["B4"].value != WORKBOOK_SCHEMA_VERSION:
            errors.append(
                f"data_product.xlsx: workbook schema metadata must be {WORKBOOK_SCHEMA_VERSION}"
            )
        if (
            workbook["_DET Metadata"]["A2"].value
            != "Compiler-owned version values. Do not edit."
        ):
            errors.append("data_product.xlsx: metadata description is stale")
        if len(str(workbook["_DET Metadata"]["B7"].value)) != 64:
            errors.append("data_product.xlsx: operator registry fingerprint is missing")
        if workbook["DET Build"]["B7"].value != "DBT_DATA_ENGINEERING_TOOLKIT_GIT_URL":
            errors.append(
                "data_product.xlsx: toolkit Git environment variable is stale"
            )
        if workbook["DET Build"]["B8"].value not in {COMPILER_VERSION, f"v{COMPILER_VERSION}"}:
            errors.append(
                f"data_product.xlsx: toolkit revision must be {COMPILER_VERSION} or v{COMPILER_VERSION}"
            )
    workbook.close()
    sample = load_workbook(
        compiler / "resources/data_product_sample.xlsx",
        read_only=True,
        data_only=True,
    )
    if sample["Fundamentals"]["C15"].value != "customer_360":
        errors.append("data_product_sample.xlsx: Customer 360 sample is missing")
    sample_quality = {
        name
        for name in sample.sheetnames
        if "quality" in name.casefold() or name.startswith("Operational ")
    }
    if sample_quality != {
        "Quality",
        "Operational Validation",
        "Operational Parameters",
    }:
        errors.append("data_product_sample.xlsx: quality surfaces are not canonical")
    sample.close()
    multi_source = load_workbook(
        compiler / "resources/data_product_multi_source_sample.xlsx",
        read_only=True,
        data_only=True,
    )
    if multi_source["DET Sources"]["A5"].value != "raw_accounts":
        errors.append("data_product_multi_source_sample.xlsx: second source is missing")
    multi_source_quality = {
        name
        for name in multi_source.sheetnames
        if "quality" in name.casefold() or name.startswith("Operational ")
    }
    if multi_source_quality != {
        "Quality",
        "Operational Validation",
        "Operational Parameters",
    }:
        errors.append(
            "data_product_multi_source_sample.xlsx: quality surfaces are not canonical"
        )
    if multi_source["DET Relationships"]["G4"].value != "many-to-one":
        errors.append(
            "data_product_multi_source_sample.xlsx: cardinality fixture is missing"
        )
    multi_source.close()
    return errors
