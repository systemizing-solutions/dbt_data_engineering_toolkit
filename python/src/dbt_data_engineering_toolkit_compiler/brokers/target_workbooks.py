"""Openpyxl broker for merging imported target structures into a DET workbook."""

from __future__ import annotations

import json
from collections.abc import Mapping
from pathlib import Path
from typing import Protocol, cast

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..services.imports.structures import (
    ImportedModel,
    ImportedProduct,
    safe_field_identifier,
    safe_identifier,
)
from .template_workbooks import refresh_workbook
from .workbooks import CellValue

EXCEL_CELL_TEXT_LIMIT = 32_767


class TargetWorkbookBroker(Protocol):
    def apply(
        self,
        workbook_path: Path,
        product: ImportedProduct,
        models: list[ImportedModel],
        *,
        replace_schema: bool = False,
        identity_mappings: bool = False,
    ) -> None: ...


def _headers(sheet, row: int) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column for cell in sheet[row] if cell.value not in (None, "")
    }


def _first_empty_row(sheet, key_column: int = 1, start: int = 4) -> int:
    row = start
    while sheet.cell(row, key_column).value not in (None, ""):
        row += 1
    return row


def _schema_sheets(workbook) -> list[Worksheet]:
    return [
        cast(Worksheet, workbook[name])
        for name in workbook.sheetnames
        if name.startswith("Schema ") and name != "Schema <table_name>"
    ]


def _schema_sheet(workbook, model: ImportedModel):
    for sheet in _schema_sheets(workbook):
        if str(sheet["B5"].value or "").strip() == model.name:
            return sheet
    candidates = _schema_sheets(workbook)
    if not candidates:
        raise ValueError("workbook has no reusable Schema sheet")
    base = candidates[0]
    base_has_properties = any(
        base.cell(row, 1).value not in (None, "") for row in range(14, base.max_row + 1)
    )
    existing_models = {
        str(workbook["DET Models"].cell(row, 1).value or "").strip()
        for row in range(4, workbook["DET Models"].max_row + 1)
    }
    if not base_has_properties and str(base["B5"].value or "") not in existing_models:
        sheet = base
    else:
        sheet = workbook.copy_worksheet(base)
    title = f"Schema {model.name}"[:31]
    if title in workbook.sheetnames and workbook[title] is not sheet:
        suffix = 2
        while f"{title[:28]}_{suffix}" in workbook.sheetnames:
            suffix += 1
        title = f"{title[:28]}_{suffix}"
    sheet.title = title
    return sheet


def _apply_product(workbook, product: ImportedProduct) -> None:
    sheet = workbook["Fundamentals"]
    sheet["C7"] = (
        product.contract_id or f"urn:datacontract:{product.domain or 'domain'}:{product.product_id}"
    )
    sheet["C8"] = product.name
    sheet["C9"] = product.version
    sheet["C10"] = product.status
    if product.owner is not None:
        sheet["C12"] = product.owner
    if product.domain is not None:
        sheet["C14"] = product.domain
    sheet["C15"] = product.product_id
    if product.description is not None:
        sheet["C19"] = product.description
    workbook["DET Build"]["B5"] = product.product_id
    if product.raw_contract:
        raw = (
            workbook["_DET Raw ODCS"]
            if "_DET Raw ODCS" in workbook.sheetnames
            else workbook.create_sheet("_DET Raw ODCS")
        )
        raw["A1"] = "Original imported ODCS; compiler-managed passthrough JSON"
        payload = json.dumps(product.raw_contract, sort_keys=True, default=str)
        for row in range(2, raw.max_row + 1):
            raw.cell(row, 1).value = None
        for row, offset in enumerate(range(0, len(payload), EXCEL_CELL_TEXT_LIMIT), start=2):
            raw.cell(row, 1, payload[offset : offset + EXCEL_CELL_TEXT_LIMIT])
        raw.sheet_state = "veryHidden"
        _apply_contract_metadata(workbook, product.raw_contract)


def _cell_value(value: object) -> CellValue:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _advanced_excel_values(raw: dict[str, object] | None) -> dict[str, CellValue]:
    if not raw:
        return {}
    values: dict[str, object] = {
        "Business Name": raw.get("businessName"),
        "Example(s)": raw.get("examples"),
        "Tags": raw.get("tags"),
        "Physical Name": raw.get("physicalName"),
        "Primary Key Position": raw.get("primaryKeyPosition"),
        "Partitioned": raw.get("partitioned"),
        "Partition Key Position": raw.get("partitionKeyPosition"),
        "Encrypted Name": raw.get("encryptedName"),
        "Transform Sources": raw.get("transformSourceObjects"),
        "Transform Logic": raw.get("transformLogic"),
        "Transform Description": raw.get("transformDescription"),
        "Critical Data Element Status": raw.get("criticalDataElement"),
    }
    definitions = raw.get("authoritativeDefinitions")
    if isinstance(definitions, list) and definitions and isinstance(definitions[0], dict):
        values["Authoritative Definition URL"] = definitions[0].get("url")
        values["Authoritative Definition Type"] = definitions[0].get("type")
    options = raw.get("logicalTypeOptions")
    option_headers = {
        "maximumItems": "Maximum Items",
        "minimumItems": "Minimum Items",
        "uniqueItems": "Unique Items",
        "format": "Format",
        "minLength": "Minimum Length",
        "maxLength": "Maximum Length",
        "exclusiveMinimum": "Exclusive Minimum",
        "minimum": "Minimum",
        "exclusiveMaximum": "Exclusive Maximum",
        "maximum": "Maximum",
        "multipleOf": "Multiple Of",
        "minProperties": "Minimum Properties",
        "maxProperties": "Maximum Properties",
        "requiredProperties": "Required Properties",
        "pattern": "Pattern",
    }
    if isinstance(options, dict):
        values.update(
            {header: options.get(key) for key, header in option_headers.items() if key in options}
        )
    return {
        key: (
            ", ".join(str(item) for item in value)
            if isinstance(value, list)
            else _cell_value(value)
        )
        for key, value in values.items()
        if value is not None
    }


def _clear_rows(sheet, start: int, max_column: int) -> None:
    for row in range(start, max(sheet.max_row, 100) + 1):
        for column in range(1, max_column + 1):
            sheet.cell(row, column).value = None


def _object_mapping(value: Mapping[object, object]) -> dict[str, object]:
    return {str(key): item for key, item in value.items()}


def _payload_rows(value: object, nested_key: str | None = None) -> list[dict[str, object]]:
    if isinstance(value, Mapping) and nested_key:
        value = value.get(nested_key) or []
    if isinstance(value, Mapping):
        return [_object_mapping(value)]
    if isinstance(value, list):
        return [_object_mapping(item) for item in value if isinstance(item, Mapping)]
    return []


def _apply_contract_metadata(workbook, payload: dict[str, object]) -> None:
    """Populate common official ODCS sheets; DET mapping decisions remain separate."""

    fundamentals = workbook["Fundamentals"]
    fundamentals["C16"] = payload.get("tenant")
    description = payload.get("description") or {}
    if isinstance(description, dict):
        fundamentals["C19"] = description.get("purpose")
        fundamentals["C20"] = description.get("limitations")
        fundamentals["C21"] = description.get("usage")
    tags = payload.get("tags")
    if isinstance(tags, list):
        fundamentals["C23"] = ", ".join(str(item) for item in tags)
    elif tags is not None:
        fundamentals["C23"] = str(tags)

    relationships = workbook["Relationships"]
    _clear_rows(relationships, 5, 5)
    for row, item in enumerate(_payload_rows(payload.get("relationships")), start=5):
        values = (
            item.get("level"),
            item.get("type"),
            item.get("from"),
            item.get("to"),
            item.get("description"),
        )
        for column, value in enumerate(values, start=1):
            relationships.cell(row, column, value)

    team = workbook["Team"]
    _clear_rows(team, 5, 7)
    team_value = payload.get("team") or []
    members = (
        _payload_rows(team_value, "members")
        if isinstance(team_value, dict)
        else _payload_rows(team_value)
    )
    for row, item in enumerate(members, start=5):
        values = (
            item.get("username") or item.get("email"),
            item.get("name"),
            item.get("description"),
            item.get("role"),
            item.get("dateIn"),
            item.get("dateOut"),
            item.get("replacedByUsername"),
        )
        for column, value in enumerate(values, start=1):
            team.cell(row, column, value)

    sla = workbook["SLA"]
    _clear_rows(sla, 7, 6)
    for row, item in enumerate(_payload_rows(payload.get("slaProperties")), start=7):
        values = (
            item.get("property") or item.get("name"),
            item.get("value"),
            item.get("extendedValue"),
            item.get("unit"),
            item.get("element"),
            item.get("driver"),
        )
        for column, value in enumerate(values, start=1):
            sla.cell(row, column, value)

    for sheet_name, key, columns in (
        (
            "Roles",
            "roles",
            ("role", "description", "access", "firstLevelApprovers", "secondLevelApprovers"),
        ),
        ("Support", "support", ("channel", "url", "description", "tool", "scope", "invitationUrl")),
        ("Custom Properties", "customProperties", ("property", "value")),
    ):
        target = workbook[sheet_name]
        _clear_rows(target, 5, len(columns))
        for row, item in enumerate(_payload_rows(payload.get(key)), start=5):
            for column, field_name in enumerate(columns, start=1):
                value = item.get(field_name)
                if isinstance(value, list):
                    value = ", ".join(str(entry) for entry in value)
                target.cell(row, column, value)

    pricing = payload.get("price") or payload.get("pricing") or {}
    if isinstance(pricing, dict):
        workbook["Pricing"]["B4"] = pricing.get("amount")
        workbook["Pricing"]["B5"] = pricing.get("currency")
        workbook["Pricing"]["B6"] = pricing.get("unit")

    servers = workbook["Servers"]
    for column in range(3, servers.max_column + 1):
        for row in range(4, servers.max_row + 1):
            servers.cell(row, column).value = None
    type_rows = {
        "bigquery": {"database": 18},
        "databricks": {"host": 23, "schema": 22},
        "glue": {"account": 26, "database": 27},
        "postgres": {"host": 36, "port": 37, "database": 38, "schema": 39},
        "snowflake": {"host": 48, "port": 49, "account": 50, "database": 51, "schema": 53},
        "sqlserver": {"host": 56, "port": 57, "database": 58, "schema": 59},
        "oracle": {"host": 62, "port": 63},
        "custom": {"account": 67, "database": 69, "host": 74, "port": 77, "schema": 81},
    }
    for column, item in enumerate(_payload_rows(payload.get("servers")), start=3):
        if column > servers.max_column:
            break
        server_type = str(item.get("type") or "custom").casefold()
        servers.cell(4, column, item.get("server") or item.get("name"))
        servers.cell(5, column, item.get("environment"))
        servers.cell(6, column, item.get("description"))
        servers.cell(8, column, server_type)
        for field_name, row in type_rows.get(server_type, type_rows["custom"]).items():
            servers.cell(row, column, item.get(field_name))


QUALITY_THRESHOLDS = (
    "mustBe",
    "mustNotBe",
    "mustBeGreaterThan",
    "mustBeGreaterOrEqualTo",
    "mustBeLessThan",
    "mustBeLessOrEqualTo",
    "mustBeBetween",
    "mustNotBeBetween",
)


def _apply_quality(workbook, models: list[ImportedModel]) -> None:
    quality = workbook["Quality"]
    quality["N4"] = "Rule ID"
    quality["O4"] = "Dimension"
    quality["P4"] = "ODCS Passthrough (JSON)"
    existing = {
        (
            str(quality.cell(row, 1).value or "").strip(),
            str(quality.cell(row, 2).value or "").strip(),
            str(
                quality.cell(row, 14).value
                or quality.cell(row, 5).value
                or quality.cell(row, 6).value
                or quality.cell(row, 9).value
                or ""
            ).strip(),
        ): row
        for row in range(5, quality.max_row + 1)
        if quality.cell(row, 1).value not in (None, "")
    }
    for model in models:
        raw = model.raw_schema or {}
        groups: list[tuple[str | None, dict[str, object]]] = []
        groups.extend((None, item) for item in _payload_rows(raw.get("quality")))
        for prop in _payload_rows(raw.get("properties")):
            if isinstance(prop, dict):
                groups.extend(
                    (safe_field_identifier(str(prop.get("name") or "")), item)
                    for item in _payload_rows(prop.get("quality"))
                )
        for property_name, rule in groups:
            rule_type = str(rule.get("type") or "library")
            metric = str(rule.get("metric") or rule.get("rule") or "") or None
            query = rule.get("query")
            engine = rule.get("engine")
            threshold = next((key for key in QUALITY_THRESHOLDS if key in rule), "")
            rule_id = str(
                rule.get("id")
                or "_".join(
                    filter(
                        None,
                        (model.name, property_name, metric or str(engine or rule_type)),
                    )
                )
            )
            identity = rule_id or str(metric or query or engine or rule_type)
            key = (model.name, property_name or "", identity)
            row = existing.get(key) or _first_empty_row(quality, start=5)
            known = {
                "id",
                "type",
                "description",
                "metric",
                "rule",
                "query",
                "engine",
                "implementation",
                "dimension",
                "severity",
                "scheduler",
                "schedule",
                *QUALITY_THRESHOLDS,
            }
            passthrough = {name: value for name, value in rule.items() if name not in known}
            implementation = rule.get("implementation")
            if isinstance(implementation, (dict, list)):
                implementation = json.dumps(implementation, sort_keys=True)
            threshold_value = rule.get(threshold) if threshold else None
            if isinstance(threshold_value, (dict, list)):
                threshold_value = json.dumps(threshold_value, sort_keys=True)
            values = (
                model.name,
                property_name,
                rule_type,
                rule.get("description"),
                metric,
                query,
                threshold,
                threshold_value,
                engine,
                implementation,
                rule.get("severity"),
                rule.get("scheduler"),
                rule.get("schedule"),
                rule_id,
                rule.get("dimension"),
                json.dumps(passthrough, sort_keys=True) if passthrough else None,
            )
            for column, value in enumerate(values, start=1):
                quality.cell(row, column, value)
            existing[key] = row


def _apply_target_import(
    workbook_path: Path,
    product: ImportedProduct,
    models: list[ImportedModel],
    *,
    replace_schema: bool = False,
    identity_mappings: bool = False,
) -> None:
    """Merge imported target structures without touching existing mapping decisions."""

    workbook = load_workbook(workbook_path, keep_vba=workbook_path.suffix.casefold() == ".xlsm")
    _apply_product(workbook, product)
    model_sheet = workbook["DET Models"]
    model_rows = {
        str(model_sheet.cell(row, 1).value or "").strip(): row
        for row in range(4, model_sheet.max_row + 1)
        if model_sheet.cell(row, 1).value not in (None, "")
    }
    input_sheet = workbook["DET Model Inputs"]
    input_keys = {
        (
            str(input_sheet.cell(row, 1).value or "").strip(),
            str(input_sheet.cell(row, 3).value or "").strip(),
        )
        for row in range(4, input_sheet.max_row + 1)
        if input_sheet.cell(row, 1).value not in (None, "")
    }

    for model in models:
        sheet = _schema_sheet(workbook, model)
        sheet["B5"] = model.name
        sheet["B6"] = "table"
        sheet["B7"] = model.description
        sheet["B9"] = model.physical_name
        headers = _headers(sheet, 13)
        existing = {
            str(sheet.cell(row, headers["Property"]).value or "").strip(): row
            for row in range(14, sheet.max_row + 1)
            if sheet.cell(row, headers["Property"]).value not in (None, "")
        }
        imported_names = {item.name for item in model.columns}
        if replace_schema:
            for name, row in existing.items():
                if name not in imported_names:
                    for column in range(1, sheet.max_column + 1):
                        sheet.cell(row, column).value = None
        for item in model.columns:
            row = existing.get(item.name) or _first_empty_row(sheet, headers["Property"], 14)
            values = _advanced_excel_values(item.raw_property)
            values.update(
                {
                    "Property": item.name,
                    "Logical Type": item.logical_type,
                    "Physical Type": item.physical_type,
                    "Description": item.description,
                    "Required": not item.nullable,
                    "Unique": item.unique,
                    "Primary Key": item.name in model.grain,
                    "Classification": item.classification,
                    "DET Implementation": (
                        "audit"
                        if item.name
                        in {
                            "_det_loaded_at",
                            "_det_invocation_id",
                            "_det_source_system",
                        }
                        else "system-generated"
                        if item.name in {"_det_rejections", "_det_warnings"}
                        else "mapped"
                    ),
                }
            )
            for header, value in values.items():
                if header in headers:
                    sheet.cell(row, headers[header], value)

        existing_model_row = model_rows.get(model.name)
        row = existing_model_row or _first_empty_row(model_sheet)
        enabled = (
            model_sheet.cell(existing_model_row, 9).value
            if existing_model_row is not None
            else "Yes"
            if identity_mappings
            else "No"
        )
        values = (
            model.name,
            model.layer,
            model.materialization,
            model.description,
            ", ".join(model.grain) or None,
            None,
            None,
            "No",
            enabled,
        )
        for column, value in enumerate(values, start=1):
            model_sheet.cell(row, column, value)
        model_rows[model.name] = row

        for order, input_relation in enumerate(model.inputs, start=1):
            if (model.name, input_relation) in input_keys:
                continue
            input_row = _first_empty_row(input_sheet)
            input_sheet.cell(input_row, 1, model.name)
            input_sheet.cell(input_row, 2, order)
            input_sheet.cell(input_row, 3, input_relation)
            input_keys.add((model.name, input_relation))

    if identity_mappings:
        _add_identity_mappings(workbook, models)
    _apply_quality(workbook, models)
    workbook.save(workbook_path)
    workbook.close()
    refresh_workbook(workbook_path)


def _add_identity_mappings(workbook, models: list[ImportedModel]) -> None:
    source_sheet = workbook["DET Sources"]
    column_sheet = workbook["DET Source Schema"]
    input_sheet = workbook["DET Model Inputs"]
    mapping_sheet = workbook["DET Mapping"]
    existing_sources = {
        str(source_sheet.cell(row, 1).value or "").strip()
        for row in range(4, source_sheet.max_row + 1)
    }
    existing_columns = {
        (
            str(column_sheet.cell(row, 1).value or "").strip(),
            str(column_sheet.cell(row, 2).value or "").strip(),
        )
        for row in range(4, column_sheet.max_row + 1)
    }
    existing_inputs = {
        (
            str(input_sheet.cell(row, 1).value or "").strip(),
            str(input_sheet.cell(row, 3).value or "").strip(),
        )
        for row in range(4, input_sheet.max_row + 1)
    }
    existing_mappings = {
        (
            str(mapping_sheet.cell(row, 1).value or "").strip(),
            str(mapping_sheet.cell(row, 2).value or "").strip(),
        )
        for row in range(4, mapping_sheet.max_row + 1)
    }
    for model in models:
        relation = safe_identifier(f"{model.name}_source")
        if relation not in existing_sources:
            row = _first_empty_row(source_sheet)
            values = (
                relation,
                "imported",
                model.physical_name or model.name,
                None,
                None,
                model.description,
            )
            for column, value in enumerate(values, start=1):
                source_sheet.cell(row, column, value)
            existing_sources.add(relation)
        if (model.name, relation) not in existing_inputs:
            row = _first_empty_row(input_sheet)
            input_sheet.cell(row, 1, model.name)
            input_sheet.cell(row, 2, 1)
            input_sheet.cell(row, 3, relation)
            existing_inputs.add((model.name, relation))
        for item in model.columns:
            if (relation, item.name) not in existing_columns:
                row = _first_empty_row(column_sheet)
                values = (
                    relation,
                    item.name,
                    item.logical_type,
                    item.physical_type,
                    "Yes" if item.nullable else "No",
                    "Yes" if item.unique else "No",
                    item.description,
                )
                for column, value in enumerate(values, start=1):
                    column_sheet.cell(row, column, value)
                existing_columns.add((relation, item.name))
            if (model.name, item.name) not in existing_mappings:
                row = _first_empty_row(mapping_sheet)
                values = (model.name, item.name, relation, item.name, 1, "Copy value")
                for column, value in enumerate(values, start=1):
                    mapping_sheet.cell(row, column, value)
                existing_mappings.add((model.name, item.name))


class OpenpyxlTargetWorkbookBroker:
    """Own every openpyxl mutation used by target imports and synchronization."""

    def apply(
        self,
        workbook_path: Path,
        product: ImportedProduct,
        models: list[ImportedModel],
        *,
        replace_schema: bool = False,
        identity_mappings: bool = False,
    ) -> None:
        _apply_target_import(
            workbook_path,
            product,
            models,
            replace_schema=replace_schema,
            identity_mappings=identity_mappings,
        )
