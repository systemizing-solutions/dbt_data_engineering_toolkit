"""End-to-end regression coverage for the reference Customer 360 product."""

from __future__ import annotations

import json
import tempfile
import unittest
from importlib.resources import files
from pathlib import Path

from openpyxl import load_workbook

from dbt_data_engineering_toolkit_compiler.brokers.target_workbooks import (
    OpenpyxlTargetWorkbookBroker,
)
from dbt_data_engineering_toolkit_compiler.brokers.template_workbooks import (
    WorkbookScaffold,
    build_workbook,
)
from dbt_data_engineering_toolkit_compiler.errors import (
    GeneratedFileDriftError,
    SpecificationValidationError,
)
from dbt_data_engineering_toolkit_compiler.generation import (
    apply_changes,
    plan_changes,
    sha256_text,
)
from dbt_data_engineering_toolkit_compiler.models import (
    Cardinality,
    FailureMode,
    JoinType,
    Materialization,
    Relationship,
    SchemaImplementation,
    SchemaProperty,
    SourceColumn,
    SourceSpecification,
)
from dbt_data_engineering_toolkit_compiler.operational import Artifact
from dbt_data_engineering_toolkit_compiler.registry import OperatorRegistry
from dbt_data_engineering_toolkit_compiler.services.emissions.service import EmissionService
from dbt_data_engineering_toolkit_compiler.services.imports.structures import (
    StructureFormat,
    StructureImportService,
)
from dbt_data_engineering_toolkit_compiler.services.validation.service import (
    SpecificationValidationService,
)
from dbt_data_engineering_toolkit_compiler.services.workbooks.service import (
    WorkbookInterpretationService,
)
from dbt_data_engineering_toolkit_compiler.version import COMPILER_VERSION


def load_specification(path: Path):
    return WorkbookInterpretationService().load(path)


def validate_specification(specification, registry: OperatorRegistry | None = None) -> None:
    SpecificationValidationService(registry=registry).validate(specification)


def emit_all(specification):
    return EmissionService().emit_all(specification)


def _load(path: Path, source_format: StructureFormat, **kwargs):
    return StructureImportService().load(path, source_format, **kwargs)


def from_ddl(path: Path):
    return _load(path, StructureFormat.DDL).as_source()


def from_dbt_manifest(path: Path):
    return _load(path, StructureFormat.DBT_MANIFEST).as_source()


def targets_from_odcs(path: Path):
    return _load(path, StructureFormat.ODCS).as_target()


def targets_from_ddl(path: Path, *, product_id: str, product_name: str):
    return _load(
        path,
        StructureFormat.DDL,
        product_id=product_id,
        product_name=product_name,
    ).as_target()


def targets_from_dbt_manifest(path: Path):
    return _load(path, StructureFormat.DBT_MANIFEST).as_target()


def apply_target_import(path: Path, product, models, **kwargs) -> None:
    OpenpyxlTargetWorkbookBroker().apply(path, product, models, **kwargs)


class WorkbookCompilerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.workbook = Path(
            str(
                files("dbt_data_engineering_toolkit_compiler")
                / "resources"
                / "data_product_sample.xlsx"
            )
        )
        cls.spec = load_specification(cls.workbook)

    def test_controlled_workbook_is_valid(self) -> None:
        validate_specification(self.spec)
        self.assertEqual(self.spec.metadata.product_id, "customer_360")
        self.assertEqual(len(self.spec.models), 1)
        self.assertEqual(len(self.spec.mappings), 6)
        self.assertEqual(len(self.spec.source_columns), 6)

    def test_workbook_is_official_superset_and_controlled(self) -> None:
        workbook = load_workbook(self.workbook, read_only=False, data_only=False)
        self.assertIn("Schema stg_customers", workbook.sheetnames)
        self.assertIn("Pricing", workbook.sheetnames)
        self.assertIn("Custom Properties", workbook.sheetnames)
        self.assertEqual(workbook["_DET Lists"].sheet_state, "veryHidden")
        self.assertEqual(workbook["_DET Metadata"].sheet_state, "veryHidden")
        self.assertTrue(workbook["Fundamentals"].protection.sheet)
        self.assertTrue(workbook["DET Mapping"].protection.sheet)
        self.assertFalse(workbook["DET Mapping"]["A4"].protection.locked)
        workbook.close()

    def test_workbook_lists_are_generated_from_registry(self) -> None:
        registry = OperatorRegistry.load()
        workbook = load_workbook(self.workbook, read_only=True, data_only=True)
        labels = {
            workbook["_DET Lists"].cell(row, 1).value
            for row in range(2, workbook["_DET Lists"].max_row + 1)
            if workbook["_DET Lists"].cell(row, 1).value
        }
        self.assertEqual(labels, set(registry.labels("transformation")))
        workbook.close()

    def test_workbook_build_is_blank_and_filename_specific_by_default(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "orders_product.xlsx"
            build_workbook(path)
            copied = load_specification(path)
            validate_specification(copied)
            self.assertEqual(copied.metadata.product_id, "orders_product")
            self.assertEqual(copied.models, [])
            self.assertEqual(copied.sources, [])
            self.assertEqual(copied.mappings, [])
            self.assertEqual(copied.rules, [])

    def test_customer_sample_requires_explicit_opt_in(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "demo.xlsx"
            build_workbook(path, include_customer_sample=True)
            copied = load_specification(path)
            validate_specification(copied)
            self.assertEqual(copied.metadata.product_id, "customer_accounts")

    def test_emits_inline_sql_and_both_contracts(self) -> None:
        artifacts = {item.path.as_posix(): item.content for item in emit_all(self.spec)}
        sql = artifacts["models/staging/stg_customers.sql"]
        self.assertIn("{{ de_toolkit.clean_string(", sql)
        self.assertIn("{{ de_toolkit.mapping(", sql)
        self.assertIn("{{ de_toolkit.assertions(", sql)
        self.assertNotIn("select_cleaned", sql)
        self.assertIn("apiVersion: v3.1.0", artifacts["contracts/customer_360.odcs.yaml"])
        self.assertIn(
            f"apiVersion: det/v{COMPILER_VERSION}", artifacts["contracts/customer_360.det.yaml"]
        )
        self.assertIn("enforced: true", artifacts["models/staging/stg_customers.yml"])
        self.assertIn("templater = dbt", artifacts[".sqlfluff"])
        self.assertIn("tests/datacontract_cli/", artifacts[".sqlfluffignore"])
        self.assertIn("sqlfluff lint models tests --config .sqlfluff", artifacts["Makefile"])
        packages = artifacts["packages.yml"]
        self.assertIn("DBT_DATA_ENGINEERING_TOOLKIT_GIT_URL", packages)
        self.assertIn(
            "https://github.com/systemizing-solutions/dbt_data_engineering_toolkit.git",
            packages,
        )
        self.assertIn(f"revision: {COMPILER_VERSION}", packages)
        self.assertIn("subdirectory: aliases/de_toolkit", packages)
        self.assertNotIn("local:", packages)

    def test_final_select_does_not_publish_internal_rule_flags(self) -> None:
        artifacts = {item.path.as_posix(): item.content for item in emit_all(self.spec)}
        sql = artifacts["models/staging/stg_customers.sql"]
        final_select = sql.rsplit("\nselect\n", 1)[1]
        self.assertNotIn("_email_invalid_valid", final_select)
        self.assertNotIn("_revenue_not_positive_valid", final_select)

    def test_rejects_contract_type_mismatch(self) -> None:
        changed = self.spec.model_copy(deep=True)
        email = next(item for item in changed.schema_properties if item.name == "email")
        email.logical_type = "boolean"
        with self.assertRaises(SpecificationValidationError) as caught:
            validate_specification(changed)
        self.assertIn("pipeline produces string but Schema declares boolean", str(caught.exception))

    def test_rejects_duckdb_operational_physical_type_mismatch(self) -> None:
        changed = self.spec.model_copy(deep=True)
        rejections = next(
            item for item in changed.schema_properties if item.name == "_det_rejections"
        )
        rejections.physical_type = "varchar"
        with self.assertRaises(SpecificationValidationError) as caught:
            validate_specification(changed)
        self.assertIn("DuckDB contract type must be 'varchar[]'", str(caught.exception))

    def test_registry_rejects_type_incompatible_operation(self) -> None:
        changed = self.spec.model_copy(deep=True)
        email_source = next(item for item in changed.source_columns if item.name == "email")
        email_source.logical_type = "number"
        email_mapping = next(item for item in changed.mappings if item.target_field == "email")
        email_mapping.source_type = "number"
        with self.assertRaises(SpecificationValidationError) as caught:
            validate_specification(changed, OperatorRegistry.load())
        self.assertIn('"Clean email" requires string', str(caught.exception))

    def test_rejects_mapping_relation_outside_model_inputs(self) -> None:
        changed = self.spec.model_copy(deep=True)
        changed.sources.append(
            SourceSpecification(relation="crm_orders", source_name="raw", table_name="raw_orders")
        )
        changed.source_columns.append(
            SourceColumn(
                relation="crm_orders",
                name="email",
                logical_type="string",
                workbook_row=4,
            )
        )
        email_mapping = next(item for item in changed.mappings if item.target_field == "email")
        email_mapping.source_relation = "crm_orders"
        with self.assertRaises(SpecificationValidationError) as caught:
            validate_specification(changed)
        self.assertIn("source relation 'crm_orders' is not an input", str(caught.exception))

    def test_rejects_unknown_source_field(self) -> None:
        changed = self.spec.model_copy(deep=True)
        mapping = next(item for item in changed.mappings if item.target_field == "email")
        mapping.source_field = "emali"
        with self.assertRaises(SpecificationValidationError) as caught:
            validate_specification(changed)
        self.assertIn("crm_customers.emali is not declared", str(caught.exception))

    def test_rejects_unmapped_contract_field(self) -> None:
        changed = self.spec.model_copy(deep=True)
        changed.schema_properties.append(
            SchemaProperty(
                object_name="stg_customers",
                name="unimplemented",
                logical_type="string",
                implementation=SchemaImplementation.MAPPED,
                workbook_sheet="Schema stg_customers",
                workbook_row=25,
            )
        )
        with self.assertRaises(SpecificationValidationError) as caught:
            validate_specification(changed)
        self.assertIn("declared as mapped but has no DET Mapping", str(caught.exception))

    def test_unique_reject_row_is_blocked(self) -> None:
        changed = self.spec.model_copy(deep=True)
        rule = changed.rules[-1]
        rule.operation = "Unique"
        rule.failure = FailureMode.REJECT
        with self.assertRaises(SpecificationValidationError) as caught:
            validate_specification(changed)
        self.assertIn("Unique only supports Fail build", str(caught.exception))

    def test_unique_is_owned_by_odcs_and_never_emits_window_where(self) -> None:
        changed = self.spec.model_copy(deep=True)
        rule = changed.rules[-1]
        rule.operation = "Unique"
        rule.name = "customer_id_duplicate"
        rule.failure = FailureMode.FAIL
        with self.assertRaises(SpecificationValidationError) as caught:
            validate_specification(changed)
        self.assertIn("Unique is already owned by the ODCS schema", str(caught.exception))
        artifacts = {item.path.as_posix(): item.content for item in emit_all(self.spec)}
        all_sql = "\n".join(value for key, value in artifacts.items() if key.endswith(".sql"))
        self.assertNotIn("count(*) over", all_sql)
        self.assertNotIn("tests/stg_customers__customer_id_duplicate.sql", artifacts)

    def test_cardinality_generates_uniqueness_test(self) -> None:
        changed = self.spec.model_copy(deep=True)
        changed.sources.append(
            SourceSpecification(relation="crm_status", source_name="raw", table_name="raw_status")
        )
        changed.source_columns.append(
            SourceColumn(
                relation="crm_status",
                name="status_code",
                logical_type="string",
                unique=True,
                workbook_row=10,
            )
        )
        changed.models[0].inputs.append("crm_status")
        changed.relationships.append(
            Relationship(
                model="stg_customers",
                left_relation="crm_customers",
                right_relation="crm_status",
                join_type=JoinType.LEFT,
                left_key="status_code",
                right_key="status_code",
                cardinality=Cardinality.MANY_TO_ONE,
                workbook_row=4,
            )
        )
        validate_specification(changed)
        sources = {item.path.as_posix(): item.content for item in emit_all(changed)}[
            "models/sources.yml"
        ]
        self.assertIn("name: raw_status", sources)
        self.assertIn("data_tests:\n      - unique", sources)

    def test_many_to_many_is_blocked_by_default(self) -> None:
        changed = self.spec.model_copy(deep=True)
        changed.sources.append(
            SourceSpecification(relation="crm_status", source_name="raw", table_name="raw_status")
        )
        changed.source_columns.append(
            SourceColumn(
                relation="crm_status", name="status_code", logical_type="string", workbook_row=10
            )
        )
        changed.models[0].inputs.append("crm_status")
        changed.relationships.append(
            Relationship(
                model="stg_customers",
                left_relation="crm_customers",
                right_relation="crm_status",
                join_type=JoinType.LEFT,
                left_key="status_code",
                right_key="status_code",
                cardinality=Cardinality.MANY_TO_MANY,
                workbook_row=4,
            )
        )
        with self.assertRaises(SpecificationValidationError) as caught:
            validate_specification(changed)
        self.assertIn("many-to-many join is blocked", str(caught.exception))

    def test_incremental_is_not_advertised_before_complete(self) -> None:
        changed = self.spec.model_copy(deep=True)
        changed.models[0].materialization = Materialization.INCREMENTAL
        with self.assertRaises(SpecificationValidationError) as caught:
            validate_specification(changed)
        self.assertIn(
            "incremental materialization is intentionally unavailable", str(caught.exception)
        )

    def test_ddl_source_import(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            ddl = Path(temp) / "source.sql"
            ddl.write_text(
                "CREATE TABLE crm_people (person_id VARCHAR PRIMARY KEY, amount DECIMAL(18,2), "
                "created_at TIMESTAMP NOT NULL, note VARCHAR DEFAULT 'a,b');",
                encoding="utf-8",
            )
            sources, columns = from_ddl(ddl)
            self.assertEqual(sources[0].relation, "crm_people")
            self.assertEqual(
                [item.name for item in columns],
                ["person_id", "amount", "created_at", "note"],
            )
            self.assertEqual(columns[1].logical_type, "number")
            self.assertFalse(columns[2].nullable)

    def test_odcs_to_workbook_import_preserves_operational_prefixes(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            contract = root / "orders.odcs.yaml"
            contract.write_text(
                """apiVersion: v3.1.0
kind: DataContract
id: urn:datacontract:sales:orders
name: Orders
version: 1.0.0
status: active
domain: sales
description:
  purpose: Governed order facts
team:
  name: Sales Data
  members:
    - username: sales@example.com
      name: Sales Owner
      role: Owner
servers:
  - server: warehouse
    type: postgres
    environment: production
    host: db.example.com
    port: 5432
    database: analytics
    schema: marts
slaProperties:
  - property: freshness
    value: 2
    unit: hours
    description: Updated every two hours
roles:
  - role: orders_reader
    access: read
support:
  - channel: slack
    url: https://example.com/support
schema:
  - name: fct_orders
    physicalType: table
    quality:
      - id: orders_not_empty
        type: library
        metric: rowCount
        mustBeGreaterThan: 0
        dimension: completeness
        method: reconciliation
    properties:
      - name: order_id
        logicalType: integer
        physicalType: bigint
        required: true
        primaryKey: true
        classification: internal
      - name: _det_loaded_at
        logicalType: timestamp
        physicalType: timestamp with time zone
""",
                encoding="utf-8",
            )
            product, models = targets_from_odcs(contract)
            workbook = root / "orders.xlsx"
            build_workbook(
                workbook,
                scaffold=WorkbookScaffold(product_id="orders", name="Orders"),
            )
            apply_target_import(workbook, product, models, replace_schema=True)
            parsed = load_specification(workbook)
            self.assertEqual(parsed.metadata.version, "1.0.0")
            self.assertEqual(parsed.metadata.description, "Governed order facts")
            self.assertEqual(parsed.metadata.owner, "Sales Data")
            self.assertEqual(parsed.team[0].username, "sales@example.com")
            self.assertEqual(parsed.servers[0].server_type, "postgres")
            self.assertEqual(parsed.servers[0].database, "analytics")
            self.assertEqual(parsed.sla[0].name, "freshness")
            self.assertEqual(parsed.roles[0].role, "orders_reader")
            self.assertEqual(parsed.support[0].channel, "slack")
            self.assertEqual(parsed.quality[0].rule_id, "orders_not_empty")
            self.assertEqual(parsed.quality[0].workbook_sheet, "Quality")
            self.assertEqual(parsed.quality[0].odcs_fields, {"method": "reconciliation"})
            imported = load_workbook(workbook, read_only=True, data_only=True)
            try:
                self.assertEqual(imported["Quality"]["N5"].value, "orders_not_empty")
                self.assertEqual(
                    json.loads(imported["Quality"]["P5"].value),
                    {"method": "reconciliation"},
                )
                self.assertEqual(
                    {
                        name
                        for name in imported.sheetnames
                        if "quality" in name.casefold() or name.startswith("Operational ")
                    },
                    {
                        "Quality",
                        "Operational Validation",
                        "Operational Parameters",
                    },
                )
            finally:
                imported.close()
            self.assertFalse(parsed.models[0].enabled)
            validate_specification(parsed)
            self.assertEqual(
                [item.name for item in parsed.schema_properties],
                ["order_id", "_det_loaded_at"],
            )
            self.assertEqual(parsed.schema_properties[0].classification, "internal")
            self.assertEqual(
                parsed.schema_properties[-1].implementation,
                SchemaImplementation.AUDIT,
            )

    def test_ddl_sync_adds_fields_without_overwriting_mappings(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            first = root / "v1.sql"
            first.write_text(
                "CREATE TABLE orders (order_id BIGINT PRIMARY KEY, amount DECIMAL(18,2));",
                encoding="utf-8",
            )
            product, models = targets_from_ddl(first, product_id="orders", product_name="Orders")
            workbook = root / "orders.xlsx"
            build_workbook(
                workbook,
                scaffold=WorkbookScaffold(product_id="orders", name="Orders"),
            )
            apply_target_import(workbook, product, models, identity_mappings=True)
            second = root / "v2.sql"
            second.write_text(
                "CREATE TABLE orders (order_id BIGINT PRIMARY KEY, amount DECIMAL(18,2), status VARCHAR);",
                encoding="utf-8",
            )
            product, models = targets_from_ddl(second, product_id="orders", product_name="Orders")
            apply_target_import(workbook, product, models)
            parsed = load_specification(workbook)
            self.assertEqual(
                {item.name for item in parsed.schema_properties},
                {"order_id", "amount", "status"},
            )
            self.assertEqual(
                {item.target_field for item in parsed.mappings},
                {"order_id", "amount"},
            )
            self.assertTrue(parsed.models[0].enabled)

    def test_dbt_manifest_imports_root_models_sources_and_dependencies(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            manifest = Path(temp) / "manifest.json"
            manifest.write_text(
                json.dumps(
                    {
                        "metadata": {"project_name": "upstream_sales"},
                        "sources": {
                            "source.upstream_sales.raw.orders": {
                                "source_name": "raw",
                                "name": "orders",
                                "identifier": "orders",
                                "columns": {
                                    "order_id": {
                                        "name": "order_id",
                                        "data_type": "bigint",
                                    }
                                },
                            }
                        },
                        "nodes": {
                            "model.upstream_sales.stg_orders": {
                                "resource_type": "model",
                                "package_name": "upstream_sales",
                                "name": "stg_orders",
                                "alias": "stg_orders",
                                "description": "Staged orders",
                                "config": {"materialized": "view"},
                                "columns": {
                                    "order_id": {
                                        "name": "order_id",
                                        "data_type": "bigint",
                                    }
                                },
                                "depends_on": {"nodes": ["source.upstream_sales.raw.orders"]},
                            },
                            "model.dependency.ignored": {
                                "resource_type": "model",
                                "package_name": "dependency",
                                "name": "ignored",
                            },
                        },
                    }
                ),
                encoding="utf-8",
            )
            product, models = targets_from_dbt_manifest(manifest)
            sources, source_columns = from_dbt_manifest(manifest)
            self.assertEqual(product.product_id, "upstream_sales")
            self.assertEqual([item.name for item in models], ["stg_orders"])
            self.assertEqual(models[0].inputs, ["raw_orders"])
            self.assertEqual([item.relation for item in sources], ["raw_orders"])
            self.assertEqual([item.name for item in source_columns], ["order_id"])

    def test_generation_is_deterministic_and_protects_drift(self) -> None:
        artifacts = [Artifact(Path("models/example.sql"), "select 1\n")]
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            changes, manifest = plan_changes(root, artifacts)
            self.assertEqual(changes[0].action, "create")
            apply_changes(root, artifacts, changes, manifest, "workbook-hash")
            changes, _ = plan_changes(root, artifacts)
            self.assertEqual(changes[0].action, "unchanged")
            (root / "models/example.sql").write_text("hand edited\n", encoding="utf-8")
            with self.assertRaises(GeneratedFileDriftError):
                plan_changes(root, artifacts)
            forced, _ = plan_changes(root, artifacts, force=True)
            self.assertEqual(forced[0].action, "update")

    def test_hash_is_stable(self) -> None:
        self.assertEqual(sha256_text("same"), sha256_text("same"))


if __name__ == "__main__":
    unittest.main()
