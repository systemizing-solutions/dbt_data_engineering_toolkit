"""Live acceptance proof for official ODCS Excel interoperability."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from dbt_data_engineering_toolkit_compiler.brokers.datacontracts import (
    CliDataContractBroker,
)
from dbt_data_engineering_toolkit_compiler.brokers.target_workbooks import (
    OpenpyxlTargetWorkbookBroker,
)
from dbt_data_engineering_toolkit_compiler.brokers.template_workbooks import (
    WorkbookScaffold,
    build_workbook,
)
from dbt_data_engineering_toolkit_compiler.services.emissions.service import EmissionService
from dbt_data_engineering_toolkit_compiler.services.imports.structures import (
    StructureFormat,
    StructureImportService,
)
from dbt_data_engineering_toolkit_compiler.services.validation.service import (
    SpecificationValidationService,
)
from dbt_data_engineering_toolkit_compiler.services.workbooks.service import (
    WorkbookInterpretationService,
)


@pytest.mark.skipif(
    shutil.which("datacontract") is None,
    reason="Data Contract CLI is installed by the package and exercised in CI",
)
def test_official_excel_to_toolkit_excel_to_linted_odcs(tmp_path: Path) -> None:
    source_yaml = tmp_path / "source.odcs.yaml"
    source_yaml.write_text(
        """apiVersion: v3.1.0
kind: DataContract
id: urn:datacontract:sales:orders
name: Orders
version: 1.0.0
status: draft
schema:
  - name: orders
    physicalType: table
    quality:
      - id: orders_not_empty
        type: library
        metric: rowCount
        mustBeGreaterThan: 0
        dimension: completeness
    properties:
      - name: order_id
        logicalType: integer
        physicalType: bigint
        required: true
        primaryKey: true
      - name: email
        logicalType: string
        quality:
          - id: email_not_null
            type: library
            metric: nullValues
            mustBe: 0
            dimension: completeness
""",
        encoding="utf-8",
    )
    broker = CliDataContractBroker()
    official_excel = tmp_path / "official_odcs.xlsx"
    export = broker.export_excel(source_yaml, official_excel)
    assert export.succeeded, export.output

    # Data Contract CLI 1.1.3 creates the official workbook correctly but its
    # exporter currently reads DataQuality.rule instead of the ODCS 3.1
    # DataQuality.metric field. Populate the two standard "Rule (Library)"
    # cells exactly as a workbook author would; no DET sheets or extensions are
    # added to this input fixture.
    workbook = load_workbook(official_excel)
    quality = workbook["Quality"]
    assert quality["E4"].value == "Rule (Library)"
    quality["E5"] = "rowCount"
    quality["E6"] = "nullValues"
    workbook.save(official_excel)
    workbook.close()

    imported = StructureImportService(datacontract=broker).load(
        official_excel,
        StructureFormat.ODCS,
    )
    product, models = imported.as_target()
    toolkit_excel = tmp_path / "orders.xlsx"
    build_workbook(
        toolkit_excel,
        scaffold=WorkbookScaffold(product_id="orders", name="Orders"),
    )
    OpenpyxlTargetWorkbookBroker().apply(
        toolkit_excel,
        product,
        models,
        replace_schema=True,
        identity_mappings=True,
    )

    specification = WorkbookInterpretationService().load(toolkit_excel)
    SpecificationValidationService().validate(specification)
    assert {
        (rule.object_name, rule.property_name, rule.metric, rule.threshold, rule.value)
        for rule in specification.quality
    } == {
        ("orders", None, "rowCount", "mustBeGreaterThan", 0),
        ("orders", "email", "nullValues", "mustBe", 0),
    }

    emitted = next(
        artifact
        for artifact in EmissionService().emit_contracts(specification)
        if artifact.path.name.endswith(".odcs.yaml")
    )
    output_yaml = tmp_path / "roundtrip.odcs.yaml"
    output_yaml.write_text(emitted.content, encoding="utf-8")
    lint = broker.lint(output_yaml)
    assert lint.succeeded, lint.output

    payload = yaml.safe_load(emitted.content)
    schema = next(item for item in payload["schema"] if item["name"] == "orders")
    email = next(item for item in schema["properties"] if item["name"] == "email")
    assert schema["quality"][0]["metric"] == "rowCount"
    assert schema["quality"][0]["mustBeGreaterThan"] == 0
    assert email["quality"][0]["metric"] == "nullValues"
    assert email["quality"][0]["mustBe"] == 0


def test_yaml_import_preserves_advanced_and_large_passthrough_attributes(
    tmp_path: Path,
) -> None:
    source_yaml = tmp_path / "source.odcs.yaml"
    payload = {
        "apiVersion": "v3.1.0",
        "kind": "DataContract",
        "id": "urn:datacontract:sales:orders",
        "name": "Orders",
        "version": "1.0.0",
        "status": "draft",
        "x-contract-extension": "x" * 33_000,
        "slaProperties": [
            {
                "property": "freshness",
                "value": 24,
                "unit": "hours",
                "driver": "operational",
                "description": "Available within one day",
            }
        ],
        "schema": [
            {
                "name": "orders",
                "physicalType": "table",
                "x-schema-extension": {"owner": "finance"},
                "properties": [
                    {
                        "name": "order_id",
                        "businessName": "Order ID",
                        "logicalType": "string",
                        "physicalType": "varchar",
                        "required": True,
                        "primaryKey": True,
                        "transformLogic": "HASH(source.order_id)",
                        "transformDescription": "Stable source identifier",
                        "logicalTypeOptions": {"format": "uuid"},
                        "x-property-extension": {"lineage": "source.order_id"},
                    }
                ],
            }
        ],
    }
    source_yaml.write_text(yaml.safe_dump(payload, sort_keys=False), encoding="utf-8")

    imported = StructureImportService().load(source_yaml, StructureFormat.ODCS)
    product, models = imported.as_target()
    toolkit_excel = tmp_path / "orders.xlsx"
    build_workbook(
        toolkit_excel,
        scaffold=WorkbookScaffold(product_id="orders", name="Orders"),
    )
    OpenpyxlTargetWorkbookBroker().apply(toolkit_excel, product, models)

    workbook = load_workbook(toolkit_excel, read_only=True)
    schema_sheet = workbook["Schema orders"]
    headers = {
        str(cell.value).strip(): cell.column
        for cell in schema_sheet[13]
        if cell.value is not None
    }
    assert schema_sheet.cell(14, headers["Transform Logic"]).value == "HASH(source.order_id)"
    assert schema_sheet.cell(14, headers["Transform Description"]).value == (
        "Stable source identifier"
    )
    assert schema_sheet.cell(14, headers["Format"]).value == "uuid"
    assert workbook["SLA"]["F7"].value == "operational"
    raw_chunks = [
        row[0].value
        for row in workbook["_DET Raw ODCS"].iter_rows(min_row=2)
        if isinstance(row[0].value, str)
    ]
    assert len(raw_chunks) > 1
    assert yaml.safe_load(source_yaml.read_text(encoding="utf-8")) == json.loads(
        "".join(raw_chunks)
    )
    workbook.close()

    specification = WorkbookInterpretationService().load(toolkit_excel)
    passthrough = specification.odcs_passthrough
    schema = passthrough["schema"]
    assert isinstance(schema, list)
    schema_item = schema[0]
    assert isinstance(schema_item, dict)
    properties = schema_item["properties"]
    assert isinstance(properties, list)
    order_id = properties[0]
    assert isinstance(order_id, dict)
    assert passthrough["x-contract-extension"] == payload["x-contract-extension"]
    assert schema_item["x-schema-extension"] == {"owner": "finance"}
    assert order_id["x-property-extension"] == {"lineage": "source.order_id"}
    assert order_id["transformLogic"] == "HASH(source.order_id)"
    sla = passthrough["slaProperties"]
    assert isinstance(sla, list)
    assert isinstance(sla[0], dict)
    assert sla[0]["description"] == "Available within one day"
