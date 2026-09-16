"""Realistic multi-source product acceptance fixture."""

from __future__ import annotations

import shutil
from importlib.resources import files
from pathlib import Path

import yaml
from openpyxl import load_workbook

from dbt_data_engineering_toolkit_compiler.brokers.template_workbooks import refresh_workbook
from dbt_data_engineering_toolkit_compiler.services.emissions.service import EmissionService
from dbt_data_engineering_toolkit_compiler.services.validation.service import (
    SpecificationValidationService,
)
from dbt_data_engineering_toolkit_compiler.services.workbooks.service import (
    WorkbookInterpretationService,
)
from dbt_data_engineering_toolkit_compiler.version import COMPILER_VERSION


def fixture_workbook() -> Path:
    return Path(
        str(
            files("dbt_data_engineering_toolkit_compiler")
            / "resources"
            / "data_product_multi_source_sample.xlsx"
        )
    )


def load_specification(path: Path):
    return WorkbookInterpretationService().load(path)


def validate_specification(specification) -> None:
    SpecificationValidationService().validate(specification)


def emit_all(specification):
    return EmissionService().emit_all(specification)


def test_multi_source_workbook_compiles_joined_readable_sql() -> None:
    spec = load_specification(fixture_workbook())
    validate_specification(spec)
    assert len(spec.sources) == 2
    assert len(spec.models) == 2
    assert len(spec.relationships) == 1
    assert len(spec.lookups) == 2

    artifacts = {item.path.as_posix(): item.content for item in emit_all(spec)}
    staging = artifacts["models/staging/stg_customers.sql"]
    mart = artifacts["models/marts/customer_account_mart.sql"]
    sources = artifacts["models/sources.yml"]
    assert "{{ de_toolkit.mapping(" in staging
    assert "{{ de_toolkit.string_title('customer_name') }}" in staging
    assert "transformed_02 as" in staging
    assert "from source__stg_customers as stg_customers" in mart
    assert "left join source__raw_accounts as raw_accounts" in mart
    assert "from joined" not in mart
    assert "- unique" in sources


def test_convert_value_is_available_as_an_early_typed_workbook_step() -> None:
    specification = load_specification(fixture_workbook()).model_copy(deep=True)
    mapping = next(item for item in specification.mappings if item.target_field == "lifetime_value")
    mapping.steps[0].operation = "convert_value"
    mapping.steps[0].parameters = {
        "data_type": "numeric",
        "precision": 18,
        "scale": 2,
    }

    validate_specification(specification)
    model_sql = {
        item.path.as_posix(): item.content for item in emit_all(specification)
    }["models/marts/customer_account_mart.sql"]

    assert "de_toolkit.convert_value(" in model_sql
    assert "'name': 'numeric'" in model_sql
    assert "'precision': 18" in model_sql
    assert "'scale': 2" in model_sql


def test_scalar_contract_transform_is_emitted_before_mapping_steps_and_preserved() -> None:
    specification = load_specification(fixture_workbook()).model_copy(deep=True)
    prop = next(item for item in specification.schema_properties if item.name == "customer_name")
    prop.odcs_fields["transformLogic"] = "upper(customer_name)"
    prop.odcs_fields["transformDescription"] = "Canonical source capitalization"

    validate_specification(specification)
    artifacts = {item.path.as_posix(): item.content for item in emit_all(specification)}
    model_sql = artifacts["models/staging/stg_customers.sql"]
    contract = yaml.safe_load(artifacts["contracts/customer_accounts.odcs.yaml"])
    contract_property = next(
        item
        for schema in contract["schema"]
        if schema["name"] == "stg_customers"
        for item in schema["properties"]
        if item["name"] == "customer_name"
    )

    assert "upper(customer_name) as customer_name" in model_sql
    assert model_sql.index("upper(customer_name) as customer_name") < model_sql.index(
        "de_toolkit.clean_string"
    )
    assert contract_property["transformLogic"] == "upper(customer_name)"
    assert contract_property["transformDescription"] == "Canonical source capitalization"


def test_workbook_has_contextual_lists_and_release_metadata() -> None:
    workbook = load_workbook(fixture_workbook(), data_only=False)
    try:
        assert workbook["_DET Metadata"]["B4"].value == COMPILER_VERSION
        assert len(str(workbook["_DET Metadata"]["B7"].value)) == 64
        assert workbook["_DET Metadata"]["B8"].value == "3.1.0"
        for name in ("_DET Lists", "_DET Metadata", "_DET Context", "_DET Raw ODCS"):
            assert workbook[name].sheet_state == "veryHidden"
            assert workbook[name].protection.sheet
        assert workbook["Quality"].sheet_state == "visible"
        assert {
            name
            for name in workbook.sheetnames
            if "quality" in name.casefold() or name.startswith("Operational ")
        } == {
            "Quality",
            "Operational Validation",
            "Operational Parameters",
        }
        assert workbook["Quality"]["N4"].value == "Rule ID"
        assert workbook["Quality"]["P4"].value == "ODCS Passthrough (JSON)"
        assert "_DET_MODEL_customer_account_mart" in workbook.defined_names
        assert "_DET_INPUTS_customer_account_mart" in workbook.defined_names
        assert "_DET_REL_raw_accounts" in workbook.defined_names
        mapping_formulas = [
            item.formula1 for item in workbook["DET Mapping"].data_validations.dataValidation
        ]
        relationship_formulas = [
            item.formula1 for item in workbook["DET Relationships"].data_validations.dataValidation
        ]
        quality_formulas = [
            item.formula1 for item in workbook["Quality"].data_validations.dataValidation
        ]
        assert any("_DET_MODEL_" in formula for formula in mapping_formulas)
        assert any("_DET_INPUTS_" in formula for formula in relationship_formulas)
        assert any("_DET_REL_" in formula for formula in relationship_formulas)
        assert any("_DET_MODEL_" in formula for formula in quality_formulas)
        schema = workbook["Schema customer_account_mart"]
        assert not schema.column_dimensions["N"].hidden
        assert schema.column_dimensions["O"].hidden
        assert not schema.column_dimensions["AM"].hidden
        assert not schema.column_dimensions["T"].hidden
        assert not schema.column_dimensions["U"].hidden
        assert "one scalar SQL expression" in schema["T13"].comment.text
    finally:
        workbook.close()


def test_refreshed_workbook_guides_but_does_not_restrict_pipeline_order(tmp_path: Path) -> None:
    workbook_path = tmp_path / "product.xlsx"
    shutil.copy2(fixture_workbook(), workbook_path)

    refresh_workbook(workbook_path)

    workbook = load_workbook(workbook_path, data_only=False)
    try:
        mapping = workbook["DET Mapping"]
        assert "guidance, not a restriction" in mapping["A2"].value
        operation_validation = next(
            item
            for item in mapping.data_validations.dataValidation
            if "F4:F100" in str(item.sqref)
        )
        assert operation_validation.errorStyle is None
        assert "Alternatives are allowed" in operation_validation.prompt
        transformation_values = {
            workbook["_DET Lists"].cell(row, 1).value
            for row in range(2, workbook["_DET Lists"].max_row + 1)
        }
        assert "Convert value" in transformation_values
        comments = [mapping.cell(row, 6).comment for row in range(4, mapping.max_row + 1)]
        assert any(
            comment is not None and "mapping format" in comment.text for comment in comments
        )
    finally:
        workbook.close()
