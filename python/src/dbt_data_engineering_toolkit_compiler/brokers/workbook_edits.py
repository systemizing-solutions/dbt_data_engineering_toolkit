"""Openpyxl-backed workbook mutation broker."""

from __future__ import annotations

from pathlib import Path
from typing import Protocol, cast

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..services.imports.structures import ImportedColumn, ImportedSource
from .template_workbooks import refresh_workbook


class SourceWorkbookEditBroker(Protocol):
    def apply_sources(
        self,
        workbook_path: Path,
        sources: list[ImportedSource],
        columns: list[ImportedColumn],
        *,
        replace: bool = False,
    ) -> None: ...


class OpenpyxlWorkbookEditBroker:
    """Localizes workbook mutations that require openpyxl's object model."""

    def apply_sources(
        self,
        workbook_path: Path,
        sources: list[ImportedSource],
        columns: list[ImportedColumn],
        *,
        replace: bool = False,
    ) -> None:
        workbook = load_workbook(workbook_path, keep_vba=workbook_path.suffix.casefold() == ".xlsm")
        source_sheet = cast(Worksheet, workbook["DET Sources"])
        column_sheet = cast(Worksheet, workbook["DET Source Schema"])
        imported_relations = {item.relation for item in sources}

        if replace:
            for row in range(column_sheet.max_row, 3, -1):
                if str(column_sheet.cell(row, 1).value or "").strip() in imported_relations:
                    column_sheet.delete_rows(row)
            for row in range(source_sheet.max_row, 3, -1):
                if str(source_sheet.cell(row, 1).value or "").strip() in imported_relations:
                    source_sheet.delete_rows(row)

        source_rows = self._existing_rows(source_sheet, 1)
        for item in sources:
            row = source_rows.get(item.relation) or self._first_empty_row(source_sheet)
            values = (
                item.relation,
                item.source_name,
                item.table_name,
                item.database,
                item.schema_name,
                item.description,
            )
            for column, value in enumerate(values, start=1):
                source_sheet.cell(row, column, value)
            source_rows[item.relation] = row

        existing_columns = {
            (
                str(column_sheet.cell(row, 1).value or "").strip(),
                str(column_sheet.cell(row, 2).value or "").strip(),
            ): row
            for row in range(4, column_sheet.max_row + 1)
            if column_sheet.cell(row, 1).value not in (None, "")
        }
        for item in columns:
            key = (item.relation, item.name)
            row = existing_columns.get(key) or self._first_empty_row(column_sheet)
            values = (
                item.relation,
                item.name,
                item.logical_type,
                item.physical_type,
                "Yes" if item.nullable else "No",
                "Yes" if item.unique else "No",
                item.description,
            )
            for column, value in enumerate(values, start=1):
                column_sheet.cell(row, column, value)
            existing_columns[key] = row

        workbook.save(workbook_path)
        workbook.close()
        refresh_workbook(workbook_path)

    @staticmethod
    def _existing_rows(sheet: Worksheet, key_column: int) -> dict[str, int]:
        return {
            str(sheet.cell(row, key_column).value).strip(): row
            for row in range(4, sheet.max_row + 1)
            if sheet.cell(row, key_column).value not in (None, "")
        }

    @staticmethod
    def _first_empty_row(sheet: Worksheet, start: int = 4) -> int:
        row = start
        while sheet.cell(row, 1).value not in (None, ""):
            row += 1
        return row
