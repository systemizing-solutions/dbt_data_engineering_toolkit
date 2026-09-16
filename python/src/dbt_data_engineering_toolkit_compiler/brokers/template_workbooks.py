"""Openpyxl writer broker for controlled DET workbook templates."""

from __future__ import annotations

import re
import shutil
from dataclasses import dataclass
from importlib.resources import as_file, files
from pathlib import Path
from typing import Iterable, Protocol

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from openpyxl.styles import Protection
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from ..adapters import AdapterRegistry
from ..registry import OperatorRegistry
from ..version import (
    COMPILER_VERSION,
    DEFAULT_DATA_PRODUCT_VERSION,
    DEFAULT_TOOLKIT_REVISION,
    ODCS_VERSION,
    OPERATOR_REGISTRY_VERSION,
    WORKBOOK_SCHEMA_VERSION,
)
from ..workbook_layout import (
    OPERATIONAL_PARAMETERS_SHEET,
    OPERATIONAL_RULES_SHEET,
)

PROTECTION_PASSWORD = "det"
MAX_INPUT_ROW = 100
WORKBOOK_VERSION = WORKBOOK_SCHEMA_VERSION
PARAMETER_LABELS = {
    "trim": "Trim",
    "collapse_whitespace": "Collapse Whitespace",
    "case": "Case",
    "blank_as_null": "Blank As Null",
    "keep_plus": "Keep Plus",
    "precision": "Precision",
    "scale": "Scale",
    "remove_prefix": "Remove Prefix",
    "remove_suffix": "Remove Suffix",
    "lookup": "Lookup",
    "default": "Default",
    "preserve_unmapped": "Preserve Unmapped",
    "case_sensitive": "Case Sensitive",
    "value": "Fill Value",
    "data_type": "Data Type",
    "format_pattern": "Format Pattern",
    "format_case": "Format Case",
    "min_value": "Minimum",
    "max_value": "Maximum",
    "inclusive": "Inclusive",
    "allow_null": "Allow Null",
    "values": "Values",
    "pattern": "Pattern",
    "operator": "Operator",
    "compare_to": "Compare To",
}


@dataclass(frozen=True, slots=True)
class WorkbookScaffold:
    """Production-safe answers used by the workbook questionnaire."""

    product_id: str
    name: str
    version: str = DEFAULT_DATA_PRODUCT_VERSION
    status: str = "draft"
    domain: str | None = None
    purpose: str | None = None
    owner: str | None = None
    initial_model: str | None = None
    adapter: str = "duckdb"
    target_schema: str = "main"
    toolkit_revision: str = DEFAULT_TOOLKIT_REVISION


DET_DATA_SHEETS = {
    "DET Models": 9,
    "DET Model Inputs": 3,
    "DET Sources": 6,
    "DET Source Schema": 7,
    "DET Mapping": 6,
    "DET Parameters": 5,
    "DET Relationships": 7,
    OPERATIONAL_RULES_SHEET: 6,
    OPERATIONAL_PARAMETERS_SHEET: 4,
    "DET Lookups": 3,
}


def _set_metadata(workbook, registry: OperatorRegistry) -> None:
    metadata = workbook["_DET Metadata"]
    metadata["A2"] = "Compiler-owned version values. Do not edit."
    values = (
        ("template_version", WORKBOOK_VERSION),
        ("compiler_version", COMPILER_VERSION),
        ("operator_registry_version", OPERATOR_REGISTRY_VERSION),
        ("operator_registry_sha256", registry.fingerprint()),
        ("odcs_version", ODCS_VERSION),
    )
    for row in range(4, max(metadata.max_row, 12) + 1):
        metadata.cell(row, 1).value = None
        metadata.cell(row, 2).value = None
    for row, (key, value) in enumerate(values, start=4):
        metadata.cell(row, 1, key)
        metadata.cell(row, 2, value)


def _ensure_system_sheets(workbook) -> None:
    for name, header in (
        ("_DET Context", "Compiler-generated contextual dropdown ranges"),
        ("_DET Raw ODCS", "Original imported ODCS; compiler-managed passthrough JSON"),
    ):
        sheet = workbook[name] if name in workbook.sheetnames else workbook.create_sheet(name)
        sheet["A1"] = header
        sheet.protection.sheet = True
        sheet.protection.set_password(PROTECTION_PASSWORD)
        sheet.sheet_state = "veryHidden"


def _clear_values(sheet, min_row: int, max_column: int) -> None:
    for row in sheet.iter_rows(
        min_row=min_row,
        max_row=max(sheet.max_row, MAX_INPUT_ROW),
        min_col=1,
        max_col=max_column,
    ):
        for cell in row:
            cell.value = None


def _set_build_sheet(workbook, scaffold: WorkbookScaffold | None = None) -> None:
    scaffold = scaffold or WorkbookScaffold(product_id="data_product", name="Data Product")
    sheet = workbook["DET Build"]
    rows = (
        ("Adapter", scaffold.adapter, "Warehouse dialect and generated example profile adapter"),
        ("Profile", scaffold.product_id, "Generated dbt profile name"),
        ("Target Schema", scaffold.target_schema, "Target schema for local execution"),
        (
            "Toolkit Git Env",
            "DBT_DATA_ENGINEERING_TOOLKIT_GIT_URL",
            "Environment variable containing the published toolkit Git URL",
        ),
        (
            "Toolkit Revision",
            scaffold.toolkit_revision,
            "Immutable published tag or commit installed by dbt deps",
        ),
        ("Allow Many To Many", "No", "Explicit opt-in after row-explosion review"),
        (
            "Contract Test Authority",
            "datacontract",
            "ODCS schema tests are applied by datacontract dbt sync",
        ),
    )
    for row_number, values in enumerate(rows, start=4):
        for column, value in enumerate(values, start=1):
            sheet.cell(row_number, column, value)
    sheet["A2"] = (
        "Advanced project settings. Generated packages use the published Git package and a "
        "pinned revision; local paths are never emitted. Export the Toolkit Git Env variable "
        "before dbt deps so packages load correctly."
    )


def _set_instructions(workbook) -> None:
    sheet = workbook["Instructions"]
    # The upstream ODCS workbook still carries the pre-1.0 CLI syntax. Keep the
    # official layout, but make the command executable with our default
    # datacontract-cli dependency.
    sheet["B18"] = "Convert to YAML using Data Contract CLI"
    sheet["B19"] = "datacontract import excel --source odcs.xlsx --output datacontract.yaml"
    rows = (
        (f"dbt_data_engineering_toolkit v{COMPILER_VERSION} workflow", None, None, None),
        ("Step", "Business action", "Outcome", "Command"),
        (
            "1",
            "Run the guided, safe workbook builder",
            "Product basics; no demo rows by default",
            "det workbook build PRODUCT.xlsx",
        ),
        (
            "2",
            "Import or define target structures",
            "ODCS model and property rows",
            "det workbook import ...",
        ),
        (
            "3",
            "Register/import upstream sources",
            "Checked source fields and types",
            "det source import ...",
        ),
        (
            "4",
            "Define model inputs, mappings, and joins",
            "Context-aware fields and executable cardinality checks",
            "det workbook refresh PRODUCT.xlsx",
        ),
        (
            "5",
            "Add contract rules on Quality; runtime handling on Operational Validation",
            "ODCS tests, warnings, and quarantine with one clear authority",
            None,
        ),
        (
            "6",
            "Validate the workbook",
            "Actionable cross-sheet diagnostics",
            "det validate PRODUCT.xlsx",
        ),
        (
            "7",
            "Preview generated changes",
            "No files changed",
            "det generate PRODUCT.xlsx --project-dir dbt_product --dry-run --prune",
        ),
        (
            "8",
            "Generate, synchronize, and lint",
            "dbt package refs use the public toolkit URL by default",
            "det generate ...; det check ...; optionally override DBT_DATA_ENGINEERING_TOOLKIT_GIT_URL for a fork or mirror",
        ),
        (
            "9",
            "Run isolated acceptance",
            "Contract, dbt tests, evaluator",
            "det prove PRODUCT.xlsx --project-dir dbt_product",
        ),
        (
            "10",
            "Merge later schema changes",
            "Mappings and rules preserved",
            "det workbook sync ...",
        ),
        (
            "Demo",
            "Opt in only when explicitly wanted",
            "Customer 360 training rows",
            "det workbook build demo.xlsx --sample-customer-data",
        ),
    )
    for row_number in range(31, 44):
        for column in range(1, 5):
            cell = sheet.cell(row_number, column)
            if not isinstance(cell, MergedCell):
                cell.value = None
    for row_number, values in enumerate(rows, start=31):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column)
            if not isinstance(cell, MergedCell):
                cell.value = value


def _make_blank(workbook, scaffold: WorkbookScaffold) -> None:
    """Remove every demo row and apply only questionnaire answers."""

    _ensure_system_sheets(workbook)

    fundamentals = workbook["Fundamentals"]
    for row in (7, 8, 9, 10, 12, 14, 15, 16, 19, 20, 21, 23):
        fundamentals.cell(row, 3).value = None
    namespace = scaffold.domain or "domain"
    fundamentals["C7"] = f"urn:datacontract:{namespace}:{scaffold.product_id}"
    fundamentals["C8"] = scaffold.name
    fundamentals["C9"] = scaffold.version
    fundamentals["C10"] = scaffold.status
    fundamentals["C12"] = scaffold.owner
    fundamentals["C14"] = scaffold.domain
    fundamentals["C15"] = scaffold.product_id
    fundamentals["C19"] = scaffold.purpose

    for sheet_name, max_column in DET_DATA_SHEETS.items():
        _clear_values(workbook[sheet_name], 4, max_column)

    for sheet_name, min_row, max_column in (
        ("Relationships", 5, 5),
        ("Quality", 5, 15),
        ("Support", 5, 6),
        ("Team", 5, 7),
        ("Roles", 5, 5),
        ("SLA", 7, 6),
        ("Custom Properties", 5, 2),
    ):
        _clear_values(workbook[sheet_name], min_row, max_column)
    for row in range(4, workbook["Servers"].max_row + 1):
        for column in range(3, workbook["Servers"].max_column + 1):
            workbook["Servers"].cell(row, column).value = None
    for row in range(4, 7):
        workbook["Pricing"].cell(row, 2).value = None
    workbook["SLA"]["B4"] = None

    schema_names = [
        name
        for name in workbook.sheetnames
        if name.startswith("Schema ") and name != "Schema <table_name>"
    ]
    if not schema_names:
        raise ValueError("template has no reusable Schema sheet")
    base = workbook[schema_names[0]]
    for name in schema_names[1:]:
        workbook.remove(workbook[name])
    model_name = scaffold.initial_model or "data_product"
    base.title = f"Schema {model_name}"[:31]
    for row in range(14, max(base.max_row, MAX_INPUT_ROW) + 1):
        for column in range(1, base.max_column + 1):
            base.cell(row, column).value = None
    for cell in ("B5", "B7", "B8", "B9", "B10", "B11"):
        base[cell] = None
    base["B5"] = model_name
    base["B6"] = "table"

    if scaffold.initial_model:
        model = workbook["DET Models"]
        values = (
            scaffold.initial_model,
            "staging",
            "view",
            None,
            None,
            None,
            None,
            "No",
            "Yes",
        )
        for column, value in enumerate(values, start=1):
            model.cell(4, column, value)

    _set_build_sheet(workbook, scaffold)
    _set_metadata(workbook, OperatorRegistry.load())


def _values(sheet, column: int, start_row: int, *, until_blank: bool = False) -> list[str]:
    result: list[str] = []
    for row in range(start_row, sheet.max_row + 1):
        value = sheet.cell(row, column).value
        if value is None or str(value).strip() == "":
            if until_blank and result:
                break
            continue
        result.append(str(value).strip())
    return result


def _unique(values: Iterable[str]) -> list[str]:
    return sorted({value for value in values if value})


def _schema_values(workbook) -> tuple[list[str], list[str]]:
    models: list[str] = []
    fields: list[str] = []
    for name in workbook.sheetnames:
        if not name.startswith("Schema ") or name == "Schema <table_name>":
            continue
        sheet = workbook[name]
        model = str(sheet["B5"].value or name.removeprefix("Schema ")).strip()
        models.append(model)
        fields.extend(_values(sheet, 1, 14))
    return _unique(models), _unique(fields)


def _context_token(value: str) -> str:
    token = re.sub(r"[^A-Za-z0-9_]", "_", value.strip())
    if not token or token[0].isdigit():
        token = f"value_{token}"
    return token[:180]


def _replace_defined_name(workbook, name: str, reference: str) -> None:
    if name in workbook.defined_names:
        del workbook.defined_names[name]
    workbook.defined_names.add(DefinedName(name, attr_text=reference))


def _context_lists(
    workbook,
    registry: OperatorRegistry,
    lookup_names: list[str],
) -> dict[str, str]:
    context = workbook["_DET Context"]
    for row in context.iter_rows():
        for cell in row:
            cell.value = None
    definitions: dict[str, list[str]] = {}
    model_fields: dict[str, list[str]] = {}
    relation_fields: dict[str, list[str]] = {}
    for sheet_name in workbook.sheetnames:
        if not sheet_name.startswith("Schema ") or sheet_name == "Schema <table_name>":
            continue
        sheet = workbook[sheet_name]
        model = str(sheet["B5"].value or sheet_name.removeprefix("Schema ")).strip()
        fields = _unique(_values(sheet, 1, 14))
        model_fields[model] = fields
        relation_fields[model] = fields
    for row in range(4, workbook["DET Source Schema"].max_row + 1):
        relation = str(workbook["DET Source Schema"].cell(row, 1).value or "").strip()
        field = str(workbook["DET Source Schema"].cell(row, 2).value or "").strip()
        if relation and field:
            relation_fields.setdefault(relation, []).append(field)
    model_inputs: dict[str, list[str]] = {}
    for row in range(4, workbook["DET Model Inputs"].max_row + 1):
        model = str(workbook["DET Model Inputs"].cell(row, 1).value or "").strip()
        relation = str(workbook["DET Model Inputs"].cell(row, 3).value or "").strip()
        if model and relation:
            model_inputs.setdefault(model, []).append(relation)
    for model, fields in model_fields.items():
        definitions[f"_DET_MODEL_{_context_token(model)}"] = _unique(fields)
    for model, inputs in model_inputs.items():
        definitions[f"_DET_INPUTS_{_context_token(model)}"] = _unique(inputs)
    for relation, fields in relation_fields.items():
        definitions[f"_DET_REL_{_context_token(relation)}"] = _unique(fields)
    for operator in registry.operators:
        labels = [
            PARAMETER_LABELS.get(name, name.replace("_", " ").title())
            for name in operator.parameters
        ]
        definitions[f"_DET_OP_{_context_token(operator.key)}"] = labels
        for name, parameter in operator.parameters.items():
            if parameter.values:
                definitions[f"_DET_VALUE_{_context_token(operator.key)}_{_context_token(name)}"] = [
                    str(value) for value in parameter.values
                ]
    definitions["_DET_LOOKUPS"] = lookup_names

    references: dict[str, str] = {}
    for column, (name, values) in enumerate(sorted(definitions.items()), start=1):
        context.cell(1, column, name)
        controlled = values or [""]
        for row, value in enumerate(controlled, start=2):
            context.cell(row, column, value)
        letter = get_column_letter(column)
        reference = f"{quote_sheetname(context.title)}!${letter}$2:${letter}${len(controlled) + 1}"
        _replace_defined_name(workbook, name, reference)
        references[name] = reference
    return references


def _validation(
    sheet,
    cell_range: str,
    column: str,
    length: int,
    *,
    prompt: str = "Select a supported value.",
    prompt_title: str = "Controlled input",
) -> None:
    if length <= 0:
        return
    validation = DataValidation(
        type="list",
        formula1=f"'_DET Lists'!${column}$2:${column}${length + 1}",
        allow_blank=True,
    )
    validation.error = "Choose a value from the controlled list."
    validation.errorTitle = "Unsupported value"
    validation.prompt = prompt
    validation.promptTitle = prompt_title
    validation.showErrorMessage = True
    validation.showInputMessage = True
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def _formula_validation(sheet, cell_range: str, formula: str, prompt: str) -> None:
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "Choose a value valid for the selected model, relation, or operation."
    validation.errorTitle = "Value is out of context"
    validation.prompt = prompt
    validation.promptTitle = "Context-aware input"
    validation.showErrorMessage = True
    validation.showInputMessage = True
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def _direct_named_validation(sheet, cell: str, name: str, prompt: str) -> None:
    _formula_validation(sheet, cell, f"={name}", prompt)


def _operation_parameter_validations(
    workbook,
    registry: OperatorRegistry,
    defined_names: dict[str, str],
) -> None:
    mapping_operations: dict[tuple[str, str, int], str] = {}
    mapping = workbook["DET Mapping"]
    for row in range(4, mapping.max_row + 1):
        try:
            step = int(mapping.cell(row, 5).value or 0)
        except (TypeError, ValueError):
            continue
        key = (
            str(mapping.cell(row, 1).value or "").strip(),
            str(mapping.cell(row, 2).value or "").strip(),
            step,
        )
        operation = str(mapping.cell(row, 6).value or "").strip()
        if all(key) and operation:
            mapping_operations[key] = operation
            resolved = registry.resolve(operation)
            if resolved:
                labels = [
                    PARAMETER_LABELS.get(name, name.replace("_", " ").title())
                    for name in resolved.parameters
                ]
                mapping.cell(row, 6).comment = Comment(
                    "Parameters: "
                    + (", ".join(labels) if labels else "none")
                    + "\n\nRecommended flow: convert_value/type-aware clean_* -> "
                    "standardize_*/correct_errors -> mapping -> fill_missing -> derive/conform "
                    "-> validate/route -> present. "
                    "This is guidance, not a restriction. Do not use mapping format until the "
                    "value has finished participating in semantic and business logic.",
                    "DET Compiler",
                )
    parameters = workbook["DET Parameters"]
    for row in range(4, MAX_INPUT_ROW + 1):
        try:
            step = int(parameters.cell(row, 3).value or 0)
        except (TypeError, ValueError):
            continue
        key = (
            str(parameters.cell(row, 1).value or "").strip(),
            str(parameters.cell(row, 2).value or "").strip(),
            step,
        )
        operator = registry.resolve(mapping_operations.get(key, ""))
        if operator is None:
            continue
        parameter_list = f"_DET_OP_{_context_token(operator.key)}"
        if parameter_list in defined_names:
            _direct_named_validation(
                parameters,
                f"D{row}",
                parameter_list,
                f"Parameters accepted by {operator.label}.",
            )
        label = str(parameters.cell(row, 4).value or "").strip()
        parameter_key = next(
            (
                name
                for name in operator.parameters
                if PARAMETER_LABELS.get(name, name.replace("_", " ").title()) == label
            ),
            None,
        )
        if parameter_key == "lookup" and "_DET_LOOKUPS" in defined_names:
            _direct_named_validation(
                parameters,
                f"E{row}",
                "_DET_LOOKUPS",
                "Choose a lookup declared on the Lookups sheet.",
            )
        elif parameter_key:
            values_name = (
                f"_DET_VALUE_{_context_token(operator.key)}_{_context_token(parameter_key)}"
            )
            if values_name in defined_names:
                _direct_named_validation(
                    parameters,
                    f"E{row}",
                    values_name,
                    f"Choose a supported {label} value.",
                )

    rules = workbook[OPERATIONAL_RULES_SHEET]
    rule_operations: dict[tuple[str, str], str] = {}
    for row in range(4, rules.max_row + 1):
        key = (
            str(rules.cell(row, 1).value or "").strip(),
            str(rules.cell(row, 2).value or "").strip(),
        )
        operation = str(rules.cell(row, 4).value or "").strip()
        if all(key) and operation:
            rule_operations[key] = operation
            resolved = registry.resolve(operation)
            if resolved:
                labels = [
                    PARAMETER_LABELS.get(name, name.replace("_", " ").title())
                    for name in resolved.parameters
                ]
                rules.cell(row, 4).comment = Comment(
                    "Parameters: " + (", ".join(labels) if labels else "none"),
                    "DET Compiler",
                )
    parameters = workbook[OPERATIONAL_PARAMETERS_SHEET]
    for row in range(4, MAX_INPUT_ROW + 1):
        key = (
            str(parameters.cell(row, 1).value or "").strip(),
            str(parameters.cell(row, 2).value or "").strip(),
        )
        operator = registry.resolve(rule_operations.get(key, ""))
        if operator is None:
            continue
        parameter_list = f"_DET_OP_{_context_token(operator.key)}"
        if parameter_list in defined_names:
            _direct_named_validation(
                parameters,
                f"C{row}",
                parameter_list,
                f"Parameters accepted by {operator.label}.",
            )


def _unlock(sheet, ranges: Iterable[str]) -> None:
    for cell_range in ranges:
        for row in sheet[cell_range]:
            for cell in row:
                cell.protection = Protection(locked=False)


def _protect(sheet, unlocked_ranges: Iterable[str]) -> None:
    _unlock(sheet, unlocked_ranges)
    sheet.protection.sheet = True
    sheet.protection.objects = True
    sheet.protection.scenarios = True
    sheet.protection.enable()
    sheet.protection.set_password(PROTECTION_PASSWORD)


def _configure_business_views(workbook) -> None:
    mapping = workbook["DET Mapping"]
    mapping["A1"] = "DET Mapping - ordered transformation pipeline"
    mapping["A2"] = (
        "Recommended flow (guidance, not a restriction): scalar ODCS Transform Logic -> "
        "convert_value / type-aware clean_* -> standardize_* / correct_errors -> mapping -> "
        "fill_missing -> business derivation / "
        "target conformance -> validation / routing -> final presentation. Repeat a target field "
        "with Step 1, 2, 3... Convert or clean according to target intent; do not clean every "
        "source as text. Use mapping format "
        "only after semantic and business logic is complete. det validate warns about unusual "
        "ordering but permits intentional alternatives."
    )
    quality = workbook["Quality"]
    quality["A1"] = "Quality (ODCS)"
    quality["A2"] = (
        "Author contractual data-quality rules here. These rows compile into ODCS and are "
        "executed through the Data Contract lifecycle."
    )
    quality["N4"] = "Rule ID"
    quality["O4"] = "Dimension"
    quality["P4"] = "ODCS Passthrough (JSON)"
    for column in range(14, 17):
        dimension = quality.column_dimensions[get_column_letter(column)]
        dimension.hidden = True
        dimension.outlineLevel = 1
    quality.sheet_properties.outlinePr.summaryRight = True
    rules = workbook[OPERATIONAL_RULES_SHEET]
    rules["A1"] = "Operational Validation"
    rules["A2"] = (
        "Runtime behavior: warnings, quarantine/reject routing, and deliberate non-schema build "
        "failures. ODCS remains the contract-test authority."
    )
    parameters = workbook[OPERATIONAL_PARAMETERS_SHEET]
    parameters["A1"] = "Operational Parameters"
    parameters["A2"] = "Only parameters accepted by the selected validation are relevant."

    for sheet_name in workbook.sheetnames:
        if not sheet_name.startswith("Schema ") or sheet_name == "Schema <table_name>":
            continue
        sheet = workbook[sheet_name]
        for column in range(1, 40):
            dimension = sheet.column_dimensions[get_column_letter(column)]
            dimension.hidden = False
            dimension.outlineLevel = 0
        for start, end in ((10, 13), (15, 19), (22, 38)):
            for column in range(start, end + 1):
                dimension = sheet.column_dimensions[get_column_letter(column)]
                dimension.hidden = True
                dimension.outlineLevel = 1
        sheet.cell(13, 20).comment = Comment(
            "Executable before DET Mapping steps. Enter one scalar SQL expression over fields "
            "from declared model inputs, for example UPPER(transaction_description). Do not "
            "enter SELECT, FROM, WHERE, joins, aliases, or multiple statements. det validate "
            "checks syntax using the configured adapter dialect and verifies column references.",
            "DET Compiler",
        )
        sheet.column_dimensions[get_column_letter(20)].width = 42
        sheet.column_dimensions[get_column_letter(21)].width = 32
        sheet.sheet_properties.outlinePr.summaryRight = True
        sheet.freeze_panes = "A14"


def refresh_workbook(path: Path) -> None:
    """Refresh registry lists, guided dropdowns, hidden metadata, and protection."""

    registry = OperatorRegistry.load()
    workbook = load_workbook(path, keep_vba=path.suffix.casefold() == ".xlsm")
    if "_DET Metadata" not in workbook.sheetnames:
        workbook.close()
        raise ValueError(
            f"Workbook is not a dbt_data_engineering_toolkit v{COMPILER_VERSION} workbook. "
            "Create a new workbook and import the external structure."
        )
    template_version = str(workbook["_DET Metadata"]["B4"].value or "").strip()
    if template_version != WORKBOOK_VERSION:
        workbook.close()
        raise ValueError(
            f"Workbook template version is {template_version or 'missing'}; "
            f"v{WORKBOOK_VERSION} is required. Create a new workbook and import "
            "the external structure."
        )
    _ensure_system_sheets(workbook)
    _set_instructions(workbook)
    _set_metadata(workbook, registry)
    _configure_business_views(workbook)
    build_values = {
        str(workbook["DET Build"].cell(row, 1).value or "").strip(): workbook["DET Build"]
        .cell(row, 2)
        .value
        for row in range(4, 11)
    }
    if "Toolkit Git Env" not in build_values:
        product_id = str(workbook["Fundamentals"]["C15"].value or "data_product")
        _set_build_sheet(
            workbook,
            WorkbookScaffold(
                product_id=product_id,
                name=str(workbook["Fundamentals"]["C8"].value or "Data Product"),
            ),
        )
    lists = workbook["_DET Lists"]
    for row in lists.iter_rows():
        for cell in row:
            cell.value = None

    schema_models, target_fields = _schema_values(workbook)
    model_names = _unique(_values(workbook["DET Models"], 1, 4) + schema_models)
    source_relations = _unique(_values(workbook["DET Sources"], 1, 4))
    all_relations = _unique(source_relations + model_names)
    source_fields = _unique(_values(workbook["DET Source Schema"], 2, 4) + target_fields)
    lookup_names = _unique(_values(workbook["DET Lookups"], 1, 4))
    parameters = _unique(
        PARAMETER_LABELS.get(name, name.replace("_", " ").title())
        for operator in registry.operators
        for name in operator.parameters
    )
    columns = {
        "A": ("Transformations", registry.labels("transformation")),
        "B": ("Validations", registry.labels("validation")),
        "C": ("Yes / No", ["Yes", "No"]),
        "D": ("Failure", ["allow", "warn", "reject row", "fail build"]),
        "E": ("Layer", ["staging", "intermediate", "mart"]),
        "F": ("Materialization", ["view", "table", "ephemeral"]),
        "G": ("Join Type", ["left", "inner", "right", "full"]),
        "H": (
            "Cardinality",
            ["one-to-one", "one-to-many", "many-to-one", "many-to-many"],
        ),
        "I": (
            "Logical Type",
            ["string", "number", "integer", "boolean", "date", "timestamp", "time"],
        ),
        "J": ("Case", ["preserve", "lower", "upper", "title"]),
        "K": (
            "Quality Metric",
            ["rowCount", "duplicateValues", "nullValues", "missingValues", "invalidValues"],
        ),
        "L": (
            "Threshold",
            [
                "mustBe",
                "mustNotBe",
                "mustBeGreaterThan",
                "mustBeGreaterOrEqualTo",
                "mustBeLessThan",
                "mustBeLessOrEqualTo",
                "mustBeBetween",
                "mustNotBeBetween",
            ],
        ),
        "M": (
            "Dimension",
            [
                "accuracy",
                "completeness",
                "conformity",
                "consistency",
                "coverage",
                "timeliness",
                "uniqueness",
            ],
        ),
        "N": ("Models", model_names),
        "O": ("Target Fields", target_fields),
        "P": ("Source Relations", source_relations),
        "Q": ("All Relations", all_relations),
        "R": ("Source Fields", source_fields),
        "S": ("Lookups", lookup_names),
        "T": ("Parameters", parameters),
        "U": ("Implementation", ["mapped", "audit", "system-generated"]),
        "V": ("Contract Authority", ["datacontract"]),
        "W": ("Adapters", list(AdapterRegistry.default().names())),
        "X": ("Quality Type", ["library", "sql", "custom", "text"]),
        "Y": ("Quality Severity", ["info", "warning", "error"]),
    }
    for column, (header, values) in columns.items():
        lists[f"{column}1"] = header
        for index, value in enumerate(values, start=2):
            lists[f"{column}{index}"] = value

    for name in (
        "DET Models",
        "DET Model Inputs",
        "DET Sources",
        "DET Source Schema",
        "Quality",
        "DET Mapping",
        "DET Parameters",
        "DET Relationships",
        OPERATIONAL_RULES_SHEET,
        OPERATIONAL_PARAMETERS_SHEET,
        "DET Lookups",
        "DET Build",
    ):
        workbook[name].data_validations.dataValidation = []

    _validation(workbook["DET Models"], f"B4:B{MAX_INPUT_ROW}", "E", len(columns["E"][1]))
    _validation(workbook["DET Models"], f"C4:C{MAX_INPUT_ROW}", "F", len(columns["F"][1]))
    _validation(workbook["DET Models"], f"H4:H{MAX_INPUT_ROW}", "C", len(columns["C"][1]))
    _validation(workbook["DET Models"], f"I4:I{MAX_INPUT_ROW}", "C", len(columns["C"][1]))
    _validation(workbook["DET Model Inputs"], f"A4:A{MAX_INPUT_ROW}", "N", len(model_names))
    _validation(workbook["DET Model Inputs"], f"C4:C{MAX_INPUT_ROW}", "Q", len(all_relations))
    _validation(workbook["DET Source Schema"], f"A4:A{MAX_INPUT_ROW}", "P", len(source_relations))
    _validation(workbook["DET Source Schema"], f"C4:C{MAX_INPUT_ROW}", "I", len(columns["I"][1]))
    _validation(workbook["DET Source Schema"], f"E4:F{MAX_INPUT_ROW}", "C", len(columns["C"][1]))
    defined_names = _context_lists(workbook, registry, lookup_names)
    model_formula = '=INDIRECT("_DET_MODEL_"&SUBSTITUTE(SUBSTITUTE($A4," ","_"),"-","_"))'
    relation_formula_c = '=INDIRECT("_DET_REL_"&SUBSTITUTE(SUBSTITUTE($C4," ","_"),"-","_"))'
    inputs_formula = '=INDIRECT("_DET_INPUTS_"&SUBSTITUTE(SUBSTITUTE($A4," ","_"),"-","_"))'
    official_quality = workbook["Quality"]
    _validation(official_quality, f"A5:A{MAX_INPUT_ROW}", "N", len(model_names))
    _formula_validation(
        official_quality,
        f"B5:B{MAX_INPUT_ROW}",
        '=INDIRECT("_DET_MODEL_"&SUBSTITUTE(SUBSTITUTE($A5," ","_"),"-","_"))',
        "Choose a property from the selected contract schema, or leave blank for an object rule.",
    )
    _validation(official_quality, f"C5:C{MAX_INPUT_ROW}", "X", len(columns["X"][1]))
    _validation(official_quality, f"E5:E{MAX_INPUT_ROW}", "K", len(columns["K"][1]))
    _validation(official_quality, f"G5:G{MAX_INPUT_ROW}", "L", len(columns["L"][1]))
    _validation(official_quality, f"K5:K{MAX_INPUT_ROW}", "Y", len(columns["Y"][1]))
    _validation(official_quality, f"O5:O{MAX_INPUT_ROW}", "M", len(columns["M"][1]))
    mapping = workbook["DET Mapping"]
    _validation(mapping, f"A4:A{MAX_INPUT_ROW}", "N", len(model_names))
    _formula_validation(
        mapping,
        f"B4:B{MAX_INPUT_ROW}",
        model_formula,
        "Choose a target field from the selected model.",
    )
    _formula_validation(
        mapping,
        f"C4:C{MAX_INPUT_ROW}",
        inputs_formula,
        "Choose one of the selected model's declared inputs.",
    )
    _formula_validation(
        mapping,
        f"D4:D{MAX_INPUT_ROW}",
        relation_formula_c,
        "Choose a field from the selected source relation.",
    )
    _validation(
        mapping,
        f"F4:F{MAX_INPUT_ROW}",
        "A",
        len(columns["A"][1]),
        prompt=(
            "Choose by intent. Recommended: convert/clean -> standardize/correct -> map -> "
            "default -> derive/conform -> validate/route -> present. Alternatives are allowed."
        ),
        prompt_title="Ordered transformation step",
    )
    params = workbook["DET Parameters"]
    _validation(params, f"A4:A{MAX_INPUT_ROW}", "N", len(model_names))
    _formula_validation(
        params,
        f"B4:B{MAX_INPUT_ROW}",
        model_formula,
        "Choose a target field from the selected model.",
    )
    relationships = workbook["DET Relationships"]
    _validation(relationships, f"A4:A{MAX_INPUT_ROW}", "N", len(model_names))
    _formula_validation(
        relationships,
        f"B4:C{MAX_INPUT_ROW}",
        inputs_formula,
        "Choose one of the selected model's declared inputs.",
    )
    _validation(relationships, f"D4:D{MAX_INPUT_ROW}", "G", len(columns["G"][1]))
    _formula_validation(
        relationships,
        f"E4:E{MAX_INPUT_ROW}",
        '=INDIRECT("_DET_REL_"&SUBSTITUTE(SUBSTITUTE($B4," ","_"),"-","_"))',
        "Choose a key from the left relation.",
    )
    _formula_validation(
        relationships,
        f"F4:F{MAX_INPUT_ROW}",
        '=INDIRECT("_DET_REL_"&SUBSTITUTE(SUBSTITUTE($C4," ","_"),"-","_"))',
        "Choose a key from the right relation.",
    )
    _validation(relationships, f"G4:G{MAX_INPUT_ROW}", "H", len(columns["H"][1]))
    rules = workbook[OPERATIONAL_RULES_SHEET]
    _validation(rules, f"A4:A{MAX_INPUT_ROW}", "N", len(model_names))
    _formula_validation(
        rules,
        f"C4:C{MAX_INPUT_ROW}",
        model_formula,
        "Choose a produced field from the selected model.",
    )
    _validation(rules, f"D4:D{MAX_INPUT_ROW}", "B", len(columns["B"][1]))
    _validation(rules, f"E4:E{MAX_INPUT_ROW}", "D", len(columns["D"][1]))
    rule_params = workbook[OPERATIONAL_PARAMETERS_SHEET]
    _validation(rule_params, f"A4:A{MAX_INPUT_ROW}", "N", len(model_names))
    _operation_parameter_validations(workbook, registry, defined_names)
    build = workbook["DET Build"]
    _validation(build, "B4", "W", len(columns["W"][1]))
    _validation(build, "B9", "C", len(columns["C"][1]))
    _validation(build, "B10", "V", len(columns["V"][1]))

    for sheet_name in workbook.sheetnames:
        if sheet_name.startswith("Schema ") and sheet_name != "Schema <table_name>":
            sheet = workbook[sheet_name]
            _validation(sheet, f"AM14:AM{MAX_INPUT_ROW}", "U", len(columns["U"][1]))

    editable = {
        "Fundamentals": ["C7:C10", "C12:C16", "C19:C21", "C23:C23"],
        "Relationships": [f"A5:E{MAX_INPUT_ROW}"],
        "Quality": [f"A5:P{MAX_INPUT_ROW}"],
        "Support": [f"A5:F{MAX_INPUT_ROW}"],
        "Team": [f"A5:G{MAX_INPUT_ROW}"],
        "Roles": [f"A5:E{MAX_INPUT_ROW}"],
        "SLA": ["B4:B4", f"A7:F{MAX_INPUT_ROW}"],
        "Servers": ["C4:Z85"],
        "Pricing": ["B4:B6"],
        "Custom Properties": [f"A5:B{MAX_INPUT_ROW}"],
        "DET Models": [f"A4:I{MAX_INPUT_ROW}"],
        "DET Model Inputs": [f"A4:C{MAX_INPUT_ROW}"],
        "DET Sources": [f"A4:F{MAX_INPUT_ROW}"],
        "DET Source Schema": [f"A4:G{MAX_INPUT_ROW}"],
        "DET Mapping": [f"A4:F{MAX_INPUT_ROW}"],
        "DET Parameters": [f"A4:E{MAX_INPUT_ROW}"],
        "DET Relationships": [f"A4:G{MAX_INPUT_ROW}"],
        OPERATIONAL_RULES_SHEET: [f"A4:F{MAX_INPUT_ROW}"],
        OPERATIONAL_PARAMETERS_SHEET: [f"A4:D{MAX_INPUT_ROW}"],
        "DET Lookups": [f"A4:C{MAX_INPUT_ROW}"],
        "DET Build": ["B4:B10"],
    }
    for sheet_name, ranges in editable.items():
        _protect(workbook[sheet_name], ranges)
    for sheet_name in workbook.sheetnames:
        if sheet_name.startswith("Schema ") and sheet_name != "Schema <table_name>":
            _protect(workbook[sheet_name], ["B5:B11", f"A14:AM{MAX_INPUT_ROW}"])
    workbook["Instructions"].protection.sheet = True
    workbook["_DET Lists"].sheet_state = "veryHidden"
    workbook["_DET Metadata"].sheet_state = "veryHidden"
    workbook["_DET Context"].sheet_state = "veryHidden"
    workbook["_DET Raw ODCS"].sheet_state = "veryHidden"
    for name in ("_DET Lists", "_DET Metadata", "_DET Context", "_DET Raw ODCS"):
        workbook[name].protection.sheet = True
        workbook[name].protection.set_password(PROTECTION_PASSWORD)
    workbook.save(path)
    workbook.close()


def build_workbook(
    destination: Path,
    *,
    scaffold: WorkbookScaffold | None = None,
    include_customer_sample: bool = False,
    force: bool = False,
) -> None:
    """Create a safe blank workbook, or the explicitly requested demo workbook."""

    if destination.exists() and not force:
        raise FileExistsError(destination)
    resource = (
        "data_product_multi_source_sample.xlsx" if include_customer_sample else "data_product.xlsx"
    )
    template = files("dbt_data_engineering_toolkit_compiler") / "resources" / resource
    destination.parent.mkdir(parents=True, exist_ok=True)
    with as_file(template) as source:
        shutil.copyfile(source, destination)
    if not include_customer_sample:
        scaffold = scaffold or WorkbookScaffold(
            product_id=destination.stem,
            name=destination.stem.replace("_", " ").replace("-", " ").title(),
        )
        workbook = load_workbook(
            destination, keep_vba=destination.suffix.casefold() == ".xlsm"
        )
        _make_blank(workbook, scaffold)
        workbook.save(destination)
        workbook.close()
    refresh_workbook(destination)


class TemplateWorkbookBroker(Protocol):
    def build(
        self,
        destination: Path,
        scaffold: WorkbookScaffold,
        *,
        include_customer_sample: bool = False,
        force: bool = False,
    ) -> None: ...

    def refresh(self, workbook: Path) -> None: ...


class OpenpyxlTemplateWorkbookBroker:
    """Production writer for new DET workbooks."""

    def build(
        self,
        destination: Path,
        scaffold: WorkbookScaffold,
        *,
        include_customer_sample: bool = False,
        force: bool = False,
    ) -> None:
        build_workbook(
            destination,
            scaffold=scaffold,
            include_customer_sample=include_customer_sample,
            force=force,
        )

    def refresh(self, workbook: Path) -> None:
        refresh_workbook(workbook)
