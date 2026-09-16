"""Release regression coverage for ODCS quality, refresh safety, and adapters."""

from __future__ import annotations

import shutil
from importlib.resources import files
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from dbt_data_engineering_toolkit_compiler.adapters import AdapterRegistry
from dbt_data_engineering_toolkit_compiler.brokers.template_workbooks import (
    OpenpyxlTemplateWorkbookBroker,
    WorkbookScaffold,
)
from dbt_data_engineering_toolkit_compiler.brokers.workbook_edits import (
    OpenpyxlWorkbookEditBroker,
)
from dbt_data_engineering_toolkit_compiler.errors import SpecificationValidationError
from dbt_data_engineering_toolkit_compiler.models import (
    BuildConfiguration,
    ColumnMapping,
    DataProductSpecification,
    FailureMode,
    Materialization,
    ModelLayer,
    ModelSpecification,
    ProductMetadata,
    SchemaImplementation,
    SchemaProperty,
    SourceColumn,
    SourceSpecification,
    TransformationStep,
    ValidationRule,
)
from dbt_data_engineering_toolkit_compiler.services.emissions.service import EmissionService
from dbt_data_engineering_toolkit_compiler.services.imports.structures import (
    ImportedColumn,
    ImportedSource,
)
from dbt_data_engineering_toolkit_compiler.services.validation.service import (
    SpecificationValidationService,
)
from dbt_data_engineering_toolkit_compiler.services.workbooks.service import (
    WorkbookInterpretationService,
)
from dbt_data_engineering_toolkit_compiler.version import COMPILER_VERSION, WORKBOOK_SCHEMA_VERSION


def load_specification(path: Path):
    return WorkbookInterpretationService().load(path)


def validate_specification(specification) -> None:
    SpecificationValidationService().validate(specification)


def emit_all(specification):
    return EmissionService().emit_all(specification)


def _sample_workbook() -> Path:
    return Path(
        str(
            files("dbt_data_engineering_toolkit_compiler")
            / "resources"
            / "data_product_sample.xlsx"
        )
    )


def test_refresh_accepts_macro_enabled_workbook(tmp_path: Path) -> None:
    workbook_path = tmp_path / "data_product.xlsm"
    shutil.copyfile(_sample_workbook(), workbook_path)

    OpenpyxlTemplateWorkbookBroker().refresh(workbook_path)

    validate_specification(load_specification(workbook_path))


def test_build_accepts_macro_enabled_workbook_destination(tmp_path: Path) -> None:
    workbook_path = tmp_path / "new_product.xlsm"

    OpenpyxlTemplateWorkbookBroker().build(
        workbook_path,
        WorkbookScaffold(product_id="new_product", name="New Product"),
    )

    validate_specification(load_specification(workbook_path))


def test_refresh_rejects_a_noncurrent_workbook_schema(tmp_path: Path) -> None:
    workbook_path = tmp_path / "unsupported.xlsx"
    shutil.copyfile(_sample_workbook(), workbook_path)
    workbook = load_workbook(workbook_path)
    workbook["_DET Metadata"]["B4"] = "unsupported"
    workbook.save(workbook_path)
    workbook.close()

    with pytest.raises(ValueError) as raised:
        OpenpyxlTemplateWorkbookBroker().refresh(workbook_path)

    assert f"v{WORKBOOK_SCHEMA_VERSION} is required" in str(raised.value)


def test_official_quality_rows_survive_excel_ir_and_odcs(tmp_path: Path) -> None:
    workbook_path = tmp_path / "official_quality.xlsx"
    shutil.copyfile(_sample_workbook(), workbook_path)
    workbook = load_workbook(workbook_path)
    quality = workbook["Quality"]
    quality["N4"] = "Rule ID"
    quality["O4"] = "Dimension"
    quality["P4"] = "ODCS Passthrough (JSON)"
    quality["A5"] = "stg_customers"
    quality["B5"] = "email"
    quality["C5"] = "library"
    quality["D5"] = "Email must be populated."
    quality["E5"] = "nullValues"
    quality["G5"] = "mustBe"
    quality["H5"] = 0
    quality["N5"] = "email_has_no_nulls"
    quality["O5"] = "completeness"
    quality["P5"] = '{"method":"reconciliation","unit":"rows"}'
    quality["A6"] = "stg_customers"
    quality["C6"] = "custom"
    quality["D6"] = "Vendor quality policy remains executable outside DET."
    quality["I6"] = "soda"
    quality["J6"] = '{"checks":["missing_count(email) = 0"]}'
    quality["N6"] = "soda_email_policy"
    quality["P6"] = '{"businessImpact":"regulatory"}'
    workbook.save(workbook_path)
    workbook.close()

    spec = load_specification(workbook_path)
    validate_specification(spec)
    official = [rule for rule in spec.quality if rule.workbook_sheet == "Quality"]
    assert {rule.rule_id for rule in official} == {
        "email_has_no_nulls",
        "soda_email_policy",
    }

    contracts = {
        item.path.as_posix(): item.content for item in emit_all(spec) if item.path.suffix == ".yaml"
    }
    payload = yaml.safe_load(contracts["contracts/customer_360.odcs.yaml"])
    schema = next(item for item in payload["schema"] if item["name"] == "stg_customers")
    email = next(item for item in schema["properties"] if item["name"] == "email")
    assert email["quality"] == [
        {
            "method": "reconciliation",
            "unit": "rows",
            "id": "email_has_no_nulls",
            "type": "library",
            "metric": "nullValues",
            "dimension": "completeness",
            "description": "Email must be populated.",
            "mustBe": 0,
        }
    ]
    custom = next(rule for rule in schema["quality"] if rule["id"] == "soda_email_policy")
    assert custom == {
        "businessImpact": "regulatory",
        "id": "soda_email_policy",
        "type": "custom",
        "engine": "soda",
        "implementation": {"checks": ["missing_count(email) = 0"]},
        "description": "Vendor quality policy remains executable outside DET.",
    }


def test_source_replacement_preserves_mapping_and_reports_removed_field(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "source_refresh.xlsx"
    shutil.copyfile(_sample_workbook(), workbook_path)
    before = load_specification(workbook_path)
    current_source = next(item for item in before.sources if item.relation == "crm_customers")
    remaining = [
        ImportedColumn(
            relation=item.relation,
            name=item.name,
            logical_type=item.logical_type,
            physical_type=item.physical_type,
            nullable=item.nullable,
            unique=item.unique,
            description=item.description,
        )
        for item in before.source_columns
        if item.relation == "crm_customers" and item.name != "email"
    ]
    OpenpyxlWorkbookEditBroker().apply_sources(
        workbook_path,
        [
            ImportedSource(
                relation=current_source.relation,
                source_name=current_source.source_name,
                table_name=current_source.table_name,
                database=current_source.database,
                schema_name=current_source.schema_name,
                description=current_source.description,
            )
        ],
        remaining,
        replace=True,
    )

    refreshed = load_specification(workbook_path)
    assert any(
        mapping.source_relation == "crm_customers" and mapping.source_field == "email"
        for mapping in refreshed.mappings
    )
    with pytest.raises(SpecificationValidationError) as raised:
        validate_specification(refreshed)
    rendered = str(raised.value)
    assert "DET-MAP-011" in rendered
    assert "crm_customers.email is not declared" in rendered
    assert "model=stg_customers" in rendered
    assert "target=email" in rendered


def _portable_project(adapter: str) -> DataProductSpecification:
    provider = AdapterRegistry.default().get(adapter)
    assert provider is not None
    return DataProductSpecification(
        metadata=ProductMetadata(product_id="orders", name="Orders"),
        sources=[
            SourceSpecification(
                relation="raw_orders",
                source_name="raw",
                table_name="orders",
            )
        ],
        source_columns=[
            SourceColumn(
                relation="raw_orders",
                name="order_id",
                logical_type="string",
                physical_type=provider.data_type("string", None),
                nullable=False,
                workbook_row=4,
            )
        ],
        models=[
            ModelSpecification(
                name="orders",
                layer=ModelLayer.STAGING,
                materialization=Materialization.VIEW,
                inputs=["raw_orders"],
                enforce_contract=True,
            )
        ],
        schema_properties=[
            SchemaProperty(
                object_name="orders",
                name="order_id",
                logical_type="string",
                physical_type=provider.data_type("string", None),
                required=True,
                workbook_sheet="Schema orders",
                workbook_row=14,
            ),
            SchemaProperty(
                object_name="orders",
                name="_det_loaded_at",
                logical_type="timestamp",
                physical_type=provider.operational_type("_det_loaded_at"),
                implementation=SchemaImplementation.AUDIT,
                workbook_sheet="Schema orders",
                workbook_row=15,
            ),
            SchemaProperty(
                object_name="orders",
                name="_det_invocation_id",
                logical_type="string",
                physical_type=provider.operational_type("_det_invocation_id"),
                implementation=SchemaImplementation.AUDIT,
                workbook_sheet="Schema orders",
                workbook_row=16,
            ),
            SchemaProperty(
                object_name="orders",
                name="_det_warnings",
                logical_type="string",
                physical_type=provider.operational_type("_det_warnings"),
                implementation=SchemaImplementation.SYSTEM_GENERATED,
                workbook_sheet="Schema orders",
                workbook_row=17,
            ),
        ],
        mappings=[
            ColumnMapping(
                model="orders",
                target_field="order_id",
                source_relation="raw_orders",
                source_field="order_id",
                source_type="string",
                steps=[
                    TransformationStep(
                        step=1,
                        operation="copy",
                        workbook_row=4,
                    )
                ],
            )
        ],
        rules=[
            ValidationRule(
                model="orders",
                name="order_id_warning",
                target_field="order_id",
                operation="required",
                failure=FailureMode.WARN,
                workbook_row=4,
            )
        ],
        build=BuildConfiguration(adapter=adapter, toolkit_revision=f"v{COMPILER_VERSION}"),
    )


@pytest.mark.parametrize(
    ("adapter", "dialect", "profile_marker", "warning_type"),
    [
        ("postgres", "postgres", "DET_POSTGRES_HOST", "text[]"),
        ("snowflake", "snowflake", "DET_SNOWFLAKE_ACCOUNT", "array"),
    ],
)
def test_static_generated_project_for_non_duckdb_adapters(
    adapter: str,
    dialect: str,
    profile_marker: str,
    warning_type: str,
) -> None:
    spec = _portable_project(adapter)
    validate_specification(spec)
    artifacts = {item.path.as_posix(): item.content for item in emit_all(spec)}
    assert f"dialect = {dialect}" in artifacts[".sqlfluff"]
    assert profile_marker in artifacts["profiles.yml"]
    assert "target/data_product.duckdb" not in artifacts["profiles.yml"]
    assert f"data_type: {warning_type}" in artifacts["models/staging/orders.yml"]
    assert "{{ de_toolkit.assertions(" in artifacts["models/staging/orders.sql"]
    assert "local:" not in artifacts["packages.yml"]
